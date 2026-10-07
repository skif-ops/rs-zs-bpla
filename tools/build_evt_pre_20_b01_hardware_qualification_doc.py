#!/usr/bin/env python3
"""Build and audit the B01 qualification protocol DOCX."""

from __future__ import annotations

import importlib.util
from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/EVT_PRE_20_B01_HARDWARE_QUALIFICATION_PROTOCOL_REV_A.md"
OUTPUT = ROOT / "docs/EVT_PRE_20_B01_HARDWARE_QUALIFICATION_PROTOCOL_REV_A.docx"
BASE_BUILDER = ROOT / "outputs/doc_build/build_manual_rev_b.py"


def load_builder():
    spec = importlib.util.spec_from_file_location("evt_doc_builder", BASE_BUILDER)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    return module


def set_font(run, size: int) -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for key in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        fonts.set(qn(key), "Times New Roman")


def build() -> None:
    text = SOURCE.read_text(encoding="utf-8")
    for character in ("\u2014", "\u2013", "\u2011"):
        if character in text:
            raise ValueError(f"source contains forbidden dash U+{ord(character):04X}")
    builder = load_builder()
    builder.OUT_DOCX = OUTPUT
    builder.build_docx(text)
    doc = Document(OUTPUT)
    for paragraph in doc.paragraphs:
        style = paragraph.style.name if paragraph.style else ""
        if style.startswith("Heading"):
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.space_before = Pt(4)
    footer = doc.sections[0].footer.paragraphs[0]
    footer.clear()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.line_spacing = 1.0
    run = footer.add_run("Дионея EVT-PRE-20. Квалификация DIO-EVT-B01 Rev A")
    set_font(run, 14)
    doc.core_properties.title = "Протокол квалификации DIO-EVT-B01"
    doc.core_properties.subject = "Квалификация стендового образца перед решением о производственной защите"
    doc.core_properties.author = "Проект Дионея"
    doc.core_properties.keywords = "Дионея, EVT-PRE-20, DIO-EVT-B01, hardware qualification, EOL"
    doc.save(OUTPUT)


def audit() -> None:
    doc = Document(OUTPUT)
    errors: list[str] = []
    text_parts: list[str] = []
    for index, paragraph in enumerate(doc.paragraphs, 1):
        text_parts.append(paragraph.text)
        if not paragraph.text:
            continue
        style = paragraph.style.name if paragraph.style else ""
        if paragraph.paragraph_format.line_spacing != 1.0:
            errors.append(f"line-spacing:{index}")
        if style == "Title" or style.startswith("Heading"):
            if any(run.text and not run.bold for run in paragraph.runs):
                errors.append(f"heading-not-bold:{index}")
        elif re.match(r"^(?:\d+\. |• )", paragraph.text):
            pass
        else:
            indent = paragraph.paragraph_format.first_line_indent
            if indent is None or abs(indent.cm - 1.25) > 0.01:
                errors.append(f"indent:{index}")
            if paragraph.alignment != WD_ALIGN_PARAGRAPH.JUSTIFY:
                errors.append(f"alignment:{index}")
        for run in paragraph.runs:
            if run.text:
                if run.font.name != "Times New Roman":
                    errors.append(f"font:{index}:{run.font.name}")
                if run.font.size is None or round(run.font.size.pt, 1) != 14:
                    errors.append(f"size:{index}:{run.font.size}")
    for table_index, table in enumerate(doc.tables, 1):
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    text_parts.append(paragraph.text)
                    if paragraph.paragraph_format.line_spacing != 1.0:
                        errors.append(f"table-spacing:{table_index}")
                    for run in paragraph.runs:
                        if run.text and (run.font.name != "Times New Roman" or run.font.size is None or round(run.font.size.pt, 1) != 12):
                            errors.append(f"table-font:{table_index}")
    for section in doc.sections:
        for paragraph in section.footer.paragraphs:
            text_parts.append(paragraph.text)
            for run in paragraph.runs:
                if run.text and (run.font.name != "Times New Roman" or run.font.size is None or round(run.font.size.pt, 1) != 14):
                    errors.append("footer-font-or-size")
    combined = "\n".join(text_parts)
    for character in ("\u2014", "\u2013", "\u2011"):
        if character in combined:
            errors.append(f"forbidden-dash:U+{ord(character):04X}")
    if errors:
        raise RuntimeError("DOCX format audit failed: " + ", ".join(errors[:40]))


def main() -> None:
    build()
    audit()
    print(f"wrote {OUTPUT}")
    print("DOCX format audit: PASS")


if __name__ == "__main__":
    main()
