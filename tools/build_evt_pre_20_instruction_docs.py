from __future__ import annotations

import re
from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"

NAMES = [
    "EVT_PRE_20_PKI_CERTIFICATE_LIFECYCLE_INSTRUCTION_REV_A",
    "EVT_PRE_20_TIME_SYNCHRONIZATION_TRUST_INSTRUCTION_REV_A",
    "EVT_PRE_20_SERVER_DEPLOYMENT_BACKUP_OPERATIONS_INSTRUCTION_REV_A",
    "EVT_PRE_20_INSTRUCTION_REGISTER_REV_A",
]

REQUIRED = {
    NAMES[0]: ["DIO-EVT-001...040", "station-revoke", "consumer-revoke", "Резервное копирование", "offline copies"],
    NAMES[1]: ["GNSS TIME TRUSTED", "HOLDOVER", "120 s", "installation coordinates", "FIELD READY"],
    NAMES[2]: ["backup_ubuntu.sh", "restore_windows.ps1", "CRL", "audit chain", "BUILD ID"],
    NAMES[3]: ["PCB MAIN 8L", "EVT PRE 20 REV D9 FULL", "Сертификаты и PKI", "Синхронизация", "Undefined mandatory limit"],
}


def set_cell_fill(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    item = OxmlElement("w:tblHeader")
    item.set(qn("w:val"), "true")
    tr_pr.append(item)


def set_row_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_keep_with_next(paragraph, value=True) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    node = p_pr.find(qn("w:keepNext"))
    if value and node is None:
        p_pr.append(OxmlElement("w:keepNext"))
    elif not value and node is not None:
        p_pr.remove(node)


def add_page_number(paragraph) -> None:
    paragraph.add_run("Страница / Page ")
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def add_run_with_emphasis(paragraph, text: str) -> None:
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        else:
            paragraph.add_run(part)


def table_widths(cols: int) -> list[float]:
    if cols == 2:
        return [2.15, 4.85]
    if cols == 3:
        return [1.6, 3.4, 2.0]
    if cols == 4:
        return [1.35, 2.65, 1.35, 1.65]
    return [7.0 / cols] * cols


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if len(rows) < 2:
        return
    headers = rows[0]
    data = rows[2:] if all(set(cell.strip()) <= set("-: ") for cell in rows[1]) else rows[1:]
    cols = len(headers)
    widths = table_widths(cols)
    table = doc.add_table(rows=1, cols=cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.width = Inches(widths[idx])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_fill(cell, "1F4E78")
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(header.strip())
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(255, 255, 255)
    set_repeat_header(table.rows[0])
    set_row_cant_split(table.rows[0])
    for row_idx, values in enumerate(data):
        row = table.add_row()
        set_row_cant_split(row)
        cells = row.cells
        for idx in range(cols):
            value = values[idx].strip() if idx < len(values) else ""
            cell = cells[idx]
            cell.width = Inches(widths[idx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_idx % 2:
                set_cell_fill(cell, "EDF3F8")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(value)
            run.font.name = "Arial"
            run.font.size = Pt(8.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def configure(doc: Document, title: str) -> None:
    sec = doc.sections[0]
    sec.page_width = Inches(8.27)
    sec.page_height = Inches(11.69)
    sec.top_margin = Inches(0.55)
    sec.bottom_margin = Inches(0.55)
    sec.left_margin = Inches(0.7)
    sec.right_margin = Inches(0.7)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9.7)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing = 1.02
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.widow_control = True

    for name, size, before, after in (
        ("Title", 16, 0, 9),
        ("Heading 1", 14, 11, 5),
        ("Heading 2", 12, 9, 4),
        ("Heading 3", 11, 7, 3),
    ):
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        ppr = style.element.get_or_add_pPr()
        pbdr = ppr.find(qn("w:pBdr"))
        if pbdr is not None:
            ppr.remove(pbdr)

    for list_name in ("List Bullet", "List Number"):
        style = doc.styles[list_name]
        style.font.name = "Arial"
        style.font.size = Pt(10)
        style.paragraph_format.space_after = Pt(3)

    header = sec.header.paragraphs[0]
    header.clear()
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = header.add_run("Дионея EVT PRE 20")
    run.font.name = "Arial"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(89, 89, 89)

    footer = sec.footer.paragraphs[0]
    footer.clear()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("REV A   ")
    run.font.name = "Arial"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(89, 89, 89)
    add_page_number(footer)

    doc.core_properties.title = title
    doc.core_properties.subject = "EVT-PRE-20 bilingual controlled instruction"
    doc.core_properties.author = "Проект Дионея"
    doc.core_properties.keywords = "EVT-PRE-20, Dioneya, manufacturing, instruction, RU, EN"


def build(name: str) -> Path:
    source = DOCS / f"{name}.md"
    output = DOCS / f"{name}.docx"
    lines = source.read_text(encoding="utf-8").splitlines()
    doc = Document()
    configure(doc, lines[0].lstrip("# ").strip())
    index = 0
    while index < len(lines):
        line = lines[index].rstrip()
        if not line:
            index += 1
            continue
        if line.startswith("|"):
            block = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                block.append([cell.strip() for cell in lines[index].strip().strip("|").split("|")])
                index += 1
            add_table(doc, block)
            continue
        if line.startswith("# "):
            p = doc.add_paragraph(style="Title")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_run_with_emphasis(p, line[2:].strip())
            index += 1
            continue
        if line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=1)
            index += 1
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
            index += 1
            continue
        if line.startswith("    "):
            code_lines = []
            while index < len(lines) and lines[index].startswith("    "):
                code_lines.append(lines[index][4:])
                index += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.right_indent = Inches(0.1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_together = True
            p_pr = p._p.get_or_add_pPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), "F2F2F2")
            p_pr.append(shd)
            run = p.add_run("\n".join(code_lines))
            run.font.name = "Consolas"
            run.font.size = Pt(8.2)
            continue
        if re.match(r"^\d+\.\s", line):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.first_line_indent = Inches(-0.18)
            p.paragraph_format.space_after = Pt(3)
            add_run_with_emphasis(p, line)
            index += 1
            continue
        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_run_with_emphasis(p, line[2:].strip())
            index += 1
            continue
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_run_with_emphasis(p, line)
        index += 1

    doc.save(output)
    check = Document(output)
    text = "\n".join(p.text for p in check.paragraphs)
    text += "\n" + "\n".join(cell.text for table in check.tables for row in table.rows for cell in row.cells)
    for phrase in REQUIRED[name]:
        if phrase not in text:
            raise RuntimeError(f"{name}: required phrase missing: {phrase}")
    return output


for document_name in NAMES:
    print(build(document_name))
