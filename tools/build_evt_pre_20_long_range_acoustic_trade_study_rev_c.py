from __future__ import annotations

import hashlib
import math
import os
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
REPO_ROOT = ROOT.parent
REFERENCE = Path(
    os.environ.get(
        "EVT_TEMPLATE_DOCX",
        REPO_ROOT / "docs" / "EVT_PRE_20_ASSEMBLY_SETUP_INSTRUCTION_REV_B.docx",
    )
)
OUT = Path(os.environ.get("EVT_ACOUSTIC_DOC_OUTPUT", REPO_ROOT / "docs"))
OUT.mkdir(parents=True, exist_ok=True)

STEM = "EVT_PRE_20_LONG_RANGE_ACOUSTIC_HEAD_TRADE_STUDY_REV_C"
DOCX = OUT / f"{STEM}.docx"
MD = OUT / f"{STEM}.md"
FIGURES = Path(os.environ.get("EVT_ACOUSTIC_FIGURES_OUTPUT", OUT / "figures"))
FIGURES.mkdir(parents=True, exist_ok=True)
LOSS_FIG = FIGURES / "acoustic_path_loss_2km_rev_c_ru_en.png"
LAYOUT_FIG = FIGURES / "recommended_acoustic_head_layout_rev_c_ru_en.png"
EXPECTED_REFERENCE_SHA256 = "81e7338683b80e1a7facd20b548e5cad1d01bf5f32f57f13bb7e86298dacfbcd"

NAVY = "142A4A"
TABLE_BLUE = "245985"
PALE_BLUE = "EAF1F7"
LIGHT_BORDER = "B7C3CF"


def absorption_db_per_m(freq_hz: float, temp_c: float = 15.0, rh_percent: float = 70.0) -> float:
    temperature = 273.15 + temp_c
    reference_temperature = 293.15
    saturation_ratio = 10 ** (-6.8346 * (273.16 / temperature) ** 1.261 + 4.6151)
    humidity = rh_percent * saturation_ratio
    oxygen = 24.0 + 4.04e4 * humidity * (0.02 + humidity) / (0.391 + humidity)
    nitrogen = (temperature / reference_temperature) ** -0.5 * (
        9.0 + 280.0 * humidity * math.exp(-4.17 * ((temperature / reference_temperature) ** (-1.0 / 3.0) - 1.0))
    )
    frequency_sq = freq_hz**2
    classical = 1.84e-11 * (temperature / reference_temperature) ** 0.5
    molecular = (temperature / reference_temperature) ** -2.5 * (
        0.01275 * math.exp(-2239.1 / temperature) / (oxygen + frequency_sq / oxygen)
        + 0.1068 * math.exp(-3352.0 / temperature) / (nitrogen + frequency_sq / nitrogen)
    )
    return 8.686 * frequency_sq * (classical + molecular)


def propagation_loss_db(freq_hz: float, distance_m: float = 2000.0) -> float:
    return 20.0 * math.log10(distance_m) + absorption_db_per_m(freq_hz) * distance_m


def reflector_gain_db(freq_hz: float, diameter_m: float = 0.6, efficiency: float = 0.5) -> float:
    wavelength = 343.0 / freq_hz
    gain = efficiency * (math.pi * diameter_m / wavelength) ** 2
    return max(0.0, 10.0 * math.log10(gain))


def build_figures() -> None:
    regular_path = Path("C:/Windows/Fonts/arial.ttf")
    bold_path = Path("C:/Windows/Fonts/arialbd.ttf")
    font = ImageFont.truetype(str(regular_path), 23)
    small = ImageFont.truetype(str(regular_path), 18)
    tiny = ImageFont.truetype(str(regular_path), 15)
    bold = ImageFont.truetype(str(bold_path), 25)

    image = Image.new("RGB", (1500, 820), "white")
    draw = ImageDraw.Draw(image)
    left, top, right, bottom = 150, 100, 1440, 690
    x_min, x_max = math.log10(80), math.log10(5000)
    y_min, y_max = 60.0, 145.0

    def point(freq: float, value: float) -> tuple[int, int]:
        x = left + (math.log10(freq) - x_min) / (x_max - x_min) * (right - left)
        y = bottom - (value - y_min) / (y_max - y_min) * (bottom - top)
        return int(x), int(y)

    for value in range(60, 146, 10):
        _, y = point(80, value)
        draw.line((left, y, right, y), fill="#D7DEE7", width=1)
        draw.text((70, y - 10), str(value), font=small, fill="#1F2937")
    for frequency in (80, 100, 200, 500, 1000, 2000, 5000):
        x, _ = point(frequency, 60)
        draw.line((x, top, x, bottom), fill="#E5E7EB", width=1)
        label = str(frequency)
        box = draw.textbbox((0, 0), label, font=small)
        draw.text((x - (box[2] - box[0]) / 2, bottom + 12), label, font=small, fill="#1F2937")
    draw.line((left, top, left, bottom), fill="#111827", width=2)
    draw.line((left, bottom, right, bottom), fill="#111827", width=2)

    frequencies = [10 ** (x_min + index * (x_max - x_min) / 400.0) for index in range(401)]
    series = [
        ("Сферическое расхождение / Spherical spreading", "#6B7280", [20.0 * math.log10(2000.0)] * len(frequencies)),
        ("С воздухом / With air absorption", "#123B63", [propagation_loss_db(freq) for freq in frequencies]),
        ("Отражатель 0.6 m / 0.6 m reflector estimate", "#B45309", [propagation_loss_db(freq) - reflector_gain_db(freq) for freq in frequencies]),
    ]
    for _, color, values in series:
        draw.line([point(freq, value) for freq, value in zip(frequencies, values)], fill=color, width=5)
    draw.text((left, 28), "Потери на 2 км / Propagation loss at 2 km", font=bold, fill="#111827")
    draw.text((650, 750), "Частота / Frequency Hz", font=font, fill="#111827")
    draw.text((18, 70), "Потери / Loss dB", font=small, fill="#111827")
    legend_y = 120
    for label, color, _ in series:
        draw.line((left + 25, legend_y + 9, left + 90, legend_y + 9), fill=color, width=6)
        draw.text((left + 105, legend_y - 2), label, font=tiny, fill="#111827")
        legend_y += 29
    image.save(LOSS_FIG)

    image = Image.new("RGB", (1200, 1200), "white")
    draw = ImageDraw.Draw(image)
    center_x, center_y = 600, 555
    scale = 500
    lower_radius = 0.60 * scale
    upper_radius = 0.45 * scale
    draw.ellipse(
        (center_x - lower_radius, center_y - lower_radius, center_x + lower_radius, center_y + lower_radius),
        outline="#9CA3AF",
        width=3,
    )

    def arrow(angle_deg: float, radius: float, length: float, color: str) -> None:
        angle = math.radians(angle_deg)
        x = center_x + radius * math.cos(angle)
        y = center_y - radius * math.sin(angle)
        end_x = x + length * math.cos(angle)
        end_y = y - length * math.sin(angle)
        draw.line((x, y, end_x, end_y), fill=color, width=12)
        for offset in (-0.45, 0.45):
            head_angle = angle + math.pi + offset
            draw.line((end_x, end_y, end_x + 25 * math.cos(head_angle), end_y - 25 * math.sin(head_angle)), fill=color, width=10)
        draw.ellipse((x - 16, y - 16, x + 16, y + 16), fill=color)

    for angle in (0, 90, 180, 270):
        arrow(angle, lower_radius, 145, "#123B63")
    for angle in (45, 135, 225, 315):
        arrow(angle, upper_radius, 125, "#B45309")
    draw.rectangle((center_x - 18, center_y - 18, center_x + 18, center_y + 18), fill="#111827")
    draw.text((center_x + 28, center_y - 24), "мачта / mast", font=small, fill="#111827")
    draw.text((80, 45), "Предварительная головка LR-1 / Preliminary LR-1 head", font=bold, fill="#111827")
    draw.text((80, 82), "вид сверху / top view", font=font, fill="#111827")
    draw.line((85, 1040, 145, 1040), fill="#123B63", width=12)
    draw.text((165, 1026), "нижний ярус 10° / lower tier 10°", font=tiny, fill="#111827")
    draw.line((620, 1040, 680, 1040), fill="#B45309", width=12)
    draw.text((700, 1026), "верхний ярус 40-45° / upper tier 40-45°", font=tiny, fill="#111827")
    image.save(LAYOUT_FIG)


def verify_reference() -> None:
    digest = hashlib.sha256(REFERENCE.read_bytes()).hexdigest()
    if digest != EXPECTED_REFERENCE_SHA256:
        raise RuntimeError(f"Template SHA-256 changed: {digest}")


def clear_document_body(document: Document) -> None:
    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def set_run_font(run, size: float | None = None, bold: bool | None = None, color: str | None = None) -> None:
    run.font.name = "Arial"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Arial")
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Arial")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def clear_paragraph(paragraph) -> None:
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def set_header_footer(document: Document) -> None:
    section = document.sections[0]
    header = section.header.paragraphs[0]
    clear_paragraph(header)
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(header.add_run("Дионея EVT-PRE-20 | Контролируемый документ / Controlled document"), 8, False, "666666")

    footer = section.footer.paragraphs[0]
    clear_paragraph(footer)
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(footer.add_run("REV C  •  EVT_PRE_20  Страница / Page "), 8, False, "666666")
    add_page_field(footer)


def set_repeat_header(row) -> None:
    row_properties = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    row_properties.append(header)


def set_cell_fill(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 85, start: int = 95, bottom: int = 85, end: int = 95) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.find(qn("w:tcMar"))
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    properties = table._tbl.tblPr
    borders = properties.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "5")
        node.set(qn("w:color"), LIGHT_BORDER)


def add_table(document: Document, headers: list[str], rows: list[list[str]], widths_cm: list[float], font_size: float = 7.8):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header_row = table.rows[0]
    set_repeat_header(header_row)
    for index, label in enumerate(headers):
        cell = header_row.cells[index]
        cell.width = Cm(widths_cm[index])
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_fill(cell, TABLE_BLUE)
        set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        set_run_font(paragraph.add_run(label), font_size, True, "FFFFFF")
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cell = cells[index]
            cell.width = Cm(widths_cm[index])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            if row_index % 2:
                set_cell_fill(cell, PALE_BLUE)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if index else WD_ALIGN_PARAGRAPH.CENTER
            set_run_font(paragraph.add_run(str(value)), font_size)
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    return table


def add_bottom_rule(paragraph) -> None:
    properties = paragraph._p.get_or_add_pPr()
    borders = properties.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        properties.append(borders)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "9")
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), "4F81BD")
    borders.append(bottom)


def add_title(document: Document) -> None:
    title_style_properties = document.styles["Title"].element.get_or_add_pPr()
    title_style_border = title_style_properties.find(qn("w:pBdr"))
    if title_style_border is not None:
        title_style_properties.remove(title_style_border)
    title = document.add_paragraph(style="Title")
    title_properties = title._p.get_or_add_pPr()
    title_border = title_properties.find(qn("w:pBdr"))
    if title_border is not None:
        title_properties.remove(title_border)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(9)
    set_run_font(title.add_run("Альтернативный акустический тракт Дионея для дальности от 2 км"), 16, True, NAVY)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(8)
    set_run_font(subtitle.add_run("Dioneya alternative long range acoustic front end"), 11, True, NAVY)
    add_bottom_rule(subtitle)


def metadata_line(document: Document, label: str, value: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(2)
    set_run_font(paragraph.add_run(f"{label}: "), 9, True)
    set_run_font(paragraph.add_run(value), 9)


def add_bilingual_paragraph(document: Document, ru: str, en: str, bold_lead_ru: str | None = None, bold_lead_en: str | None = None) -> None:
    ru_paragraph = document.add_paragraph()
    ru_paragraph.paragraph_format.space_after = Pt(1.5)
    if bold_lead_ru and ru.startswith(bold_lead_ru):
        set_run_font(ru_paragraph.add_run("RU: " + bold_lead_ru), 9, True)
        set_run_font(ru_paragraph.add_run(ru[len(bold_lead_ru):]), 9)
    else:
        set_run_font(ru_paragraph.add_run("RU: " + ru), 9)
    en_paragraph = document.add_paragraph()
    en_paragraph.paragraph_format.space_after = Pt(4)
    if bold_lead_en and en.startswith(bold_lead_en):
        set_run_font(en_paragraph.add_run("EN: " + bold_lead_en), 9, True)
        set_run_font(en_paragraph.add_run(en[len(bold_lead_en):]), 9)
    else:
        set_run_font(en_paragraph.add_run("EN: " + en), 9)


def add_numbered_pair(document: Document, ru: str, en: str) -> None:
    paragraph = document.add_paragraph(style="List Number")
    paragraph.paragraph_format.space_after = Pt(3)
    set_run_font(paragraph.add_run(f"RU: {ru} EN: {en}"), 9)


def add_heading(document: Document, text: str, level: int = 1) -> None:
    paragraph = document.add_heading(text, level=level)
    paragraph.paragraph_format.keep_with_next = True


def add_figure(document: Document, path: Path, width_in: float, caption_ru: str, caption_en: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    shape = paragraph.add_run().add_picture(str(path), width=Inches(width_in))
    shape._inline.docPr.set("descr", f"{caption_ru} / {caption_en}")
    shape._inline.docPr.set("title", caption_en)
    caption = document.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(5)
    set_run_font(caption.add_run(f"{caption_ru} / {caption_en}"), 8, False, "555555")
    caption.runs[0].italic = True


def add_sources(document: Document) -> None:
    sources = [
        "[1] TDK InvenSense T5838 product page and DS-000383 datasheet. https://www.invensense.tdk.com/en-us/products/microphone/t5838",
        "[2] Infineon IM73D122 product data. https://www.infineon.com/part/IM73D122",
        "[3] Infineon IM73A135 datasheet. https://www.infineon.com/assets/row/public/documents/24/49/infineon-im73a135-datasheet-en.pdf",
        "[4] RODE NTG5 specifications. https://rode.com/en-us/products/ntg5",
        "[5] Sennheiser MKH 416 specifications. https://docs.cloud.sennheiser.com/en-us/mkh-416-mkh-418-s/specs-mkh-416.html",
        "[6] Jekaterynczuk G. and Piotrowski Z. Outdoor Microphone Range Tests and Spectral Analysis of UAV Acoustic Signatures for Array Development. Sensors 2025, 25, 7057. https://pmc.ncbi.nlm.nih.gov/articles/PMC12656299/",
        "[7] NASA atmospheric propagation relation and absorption discussion. https://ntrs.nasa.gov/api/citations/19750024046/downloads/19750024046.pdf",
        "[8] Meissner M. The response of Helmholtz resonator to external excitation. Archives of Acoustics. https://acoustics.ippt.pan.pl/index.php/aa/article/view/471",
        "[9] Mennitt D. J., Fristrup K. M., Notaros B. M. Characterization of gain and directivity of exponential horn receivers. JASA 142, 3257-3266. https://www.engr.colostate.edu/~notaros/Papers/JASA_Acoustic_Horns.pdf",
        "[10] Kim K. and Kim Y. Monaural sound localization based on structure-induced acoustic resonance. Sensors 2015, 15, 3872-3895. https://pubmed.ncbi.nlm.nih.gov/25668214/",
        "[11] Darras K. F. A. et al. Assembling cheap high-performance microphones for recording terrestrial wildlife. F1000Research. https://f1000research.com/articles/7-1984",
    ]
    for source in sources:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Cm(0.55)
        paragraph.paragraph_format.first_line_indent = Cm(-0.55)
        paragraph.paragraph_format.space_after = Pt(1)
        set_run_font(paragraph.add_run(source), 7.5)


def build_document() -> None:
    verify_reference()
    document = Document(REFERENCE)
    clear_document_body(document)
    set_header_footer(document)
    add_title(document)

    metadata_line(document, "Документ / Document", STEM)
    metadata_line(document, "Ревизия / Revision", "Rev C")
    metadata_line(document, "Дата / Date", "27.10.2026")
    metadata_line(document, "Базовый комплект / Baseline package", "EVT-PRE-20 REV_D10")
    metadata_line(document, "Статус / Status", "ENGINEERING EVALUATION; EVT PRODUCTION BASELINE UNCHANGED")
    metadata_line(document, "Назначение / Purpose", "выбор и проверка физической акустической головки LR-1 / selection and verification of the LR-1 physical acoustic head")

    add_heading(document, "1 Решение и область применения Decision and scope")
    add_bilingual_paragraph(
        document,
        "Главный вывод. Штатная головка из четырех TDK T5838 сохраняется в EVT-PRE-20 как малопотребляющий дежурный тракт, AAD и контрольный канал. Она не дает физического основания назначить одну гарантированную дальность 2 км для всех воздушных целей. Для проверки такой дальности нужен отдельный прототип LR-1 с большой апертурой, низким собственным шумом, секторной направленностью, ветрозащитой и калиброванной записью.",
        "Main conclusion. The standard four-channel TDK T5838 head remains in EVT-PRE-20 as the low-power watch channel, AAD and reference path. It does not provide a physical basis for one guaranteed 2 km range across all airborne targets. Verification at that range requires a separate LR-1 prototype with a large aperture, low self-noise, sector directivity, wind protection and calibrated recording.",
        "Главный вывод. ",
        "Main conclusion. ",
    )
    add_bilingual_paragraph(
        document,
        "LR-1 является параллельным экспериментальным модулем. Этот документ не изменяет PCB-MAIN, PCB-PWR, PCB-MIC, производственный BOM, прошивку, распайку или критерии приемки выпущенного комплекта EVT-PRE-20.",
        "LR-1 is a parallel experimental module. This document does not change PCB-MAIN, PCB-PWR, PCB-MIC, the manufacturing BOM, firmware, pinout or acceptance criteria of the released EVT-PRE-20 package.",
    )

    add_heading(document, "2 Расчетная база дальности Range calculation basis")
    add_bilingual_paragraph(
        document,
        "На 2 км сферическое расхождение дает 66,0 дБ потерь относительно уровня на 1 м. Поглощение воздуха добавляет частотно-зависимые потери. Расчет выполнен для 15 °C, относительной влажности 70 процентов и давления 101,325 кПа. Он не учитывает отражение от земли, растительность, здания, ветер, температурные градиенты и атмосферные теневые зоны.",
        "At 2 km, spherical spreading produces 66.0 dB of loss relative to the level at 1 m. Air absorption adds frequency-dependent loss. The calculation uses 15 °C, 70 percent relative humidity and 101.325 kPa. It excludes ground interaction, vegetation, buildings, wind, temperature gradients and atmospheric shadow zones.",
    )
    add_figure(
        document,
        LOSS_FIG,
        6.65,
        "Рисунок 1 Расчетные потери распространения на 2 км",
        "Figure 1 Calculated propagation loss at 2 km",
    )
    add_table(
        document,
        ["Частота Frequency Hz", "Поглощение Absorption dB", "Полные потери Total loss dB", "70 dB SPL at 1 m", "85 dB SPL at 1 m", "100 dB SPL at 1 m"],
        [
            ["100", "0.5", "66.5", "3.5", "18.5", "33.5"],
            ["250", "2.2", "68.3", "1.7", "16.7", "31.7"],
            ["500", "4.7", "70.7", "-0.7", "14.3", "29.3"],
            ["1000", "8.2", "74.2", "-4.2", "10.8", "25.8"],
            ["2000", "17.6", "83.6", "-13.6", "1.4", "16.4"],
            ["4000", "53.2", "119.2", "-49.2", "-34.2", "-19.2"],
        ],
        [2.6, 3.0, 3.0, 2.9, 2.9, 2.9],
        7.2,
    )
    add_bilingual_paragraph(
        document,
        "Уровни 70, 85 и 100 дБ SPL являются расчетными сценариями. Они не присвоены классам целей. Дальность надо рассчитывать по измеренному узкополосному или третьоктавному спектру конкретной цели в каждом режиме.",
        "The 70, 85 and 100 dB SPL levels are calculation scenarios. They are not assigned to target classes. Range must be calculated from the measured narrowband or one-third-octave spectrum of the specific target in each operating state.",
    )

    add_heading(document, "3 Классы целей и исходные измерения Target classes and source measurements")
    add_table(
        document,
        ["Класс Class", "Физический источник Physical source", "Режимы States", "Обязательное измерение Required measurement"],
        [
            ["Малый электрический мультикоптер\nSmall electric multirotor", "Винты, моторы, регуляторы\nPropellers, motors, controllers", "Висение, набор, пролет, удаление\nHover, climb, pass, retreat", "Калиброванный спектр на 1 м по каждому режиму\nCalibrated 1 m spectrum for each state"],
            ["Крупный электрический аппарат\nLarge electric aircraft", "Крупные винты и аэродинамический поток\nLarge propellers and aerodynamic flow", "Взлет, набор, крейсер, маневр\nTakeoff, climb, cruise, maneuver", "Спектр, тяга, ракурс и геометрия\nSpectrum, thrust, aspect and geometry"],
            ["Поршневой или роторный\nPiston or rotary combustion", "Двигатель, выхлоп, редуктор, винт или ротор\nEngine, exhaust, gearbox, propeller or rotor", "Холостой ход, нагрузка, набор, крейсер\nIdle, load, climb, cruise", "Спектр, обороты, нагрузка и конфигурация выхлопа\nSpectrum, speed, load and exhaust configuration"],
            ["Винтовой самолет\nPropeller fixed wing", "Прохождение лопастей, двигатель и планер\nBlade passage, engine and airframe", "Взлет, набор, крейсер, разворот\nTakeoff, climb, cruise, turn", "Спектр, дальность, высота, курс и скорость\nSpectrum, range, altitude, course and speed"],
            ["Турбинный или реактивный\nTurbine or jet", "Турбомашина и реактивная струя\nTurbomachinery and jet flow", "Малый газ, взлетная тяга, набор, крейсер\nIdle, takeoff thrust, climb, cruise", "Спектр и уровень для каждого режима двигателя\nSpectrum and level for each engine state"],
            ["Ложные источники\nFalse sources", "Транспорт, генераторы, механизмы, ветер, биофония\nVehicles, generators, machinery, wind, biophony", "Условия конкретной площадки\nSite-specific states", "Одновременная калиброванная запись фона\nSimultaneous calibrated background recording"],
        ],
        [3.3, 5.2, 4.6, 5.2],
        7.0,
    )
    add_bilingual_paragraph(
        document,
        "Паспортная мощность двигателя, масса или общее название БПЛА не заменяют измеренный спектр звукового давления. Критерий 2 км может быть утвержден только отдельно по классу, режиму, ракурсу и погодным условиям.",
        "Rated engine power, mass or a generic aircraft label does not replace a measured sound-pressure spectrum. A 2 km criterion can be approved only for a defined class, state, aspect and weather envelope.",
    )

    add_heading(document, "4 Ограничения штатной головки Baseline head limits")
    add_bilingual_paragraph(
        document,
        "Четыре T5838 образуют компактный массив с максимальной базой около 0,165 м. Такая геометрия полезна для согласованного захвата, оценки направления в рабочей полосе и AAD, но она не создает крупную низкочастотную апертуру. Разница между капсюлями в несколько децибел не компенсирует 66 дБ геометрических потерь и шум площадки.",
        "Four T5838 microphones form a compact array with a maximum baseline of about 0.165 m. The geometry supports coherent capture, direction estimation in its usable band and AAD, but it does not create a large low-frequency aperture. A few decibels of capsule improvement cannot offset 66 dB of geometric loss and site noise.",
    )
    add_table(
        document,
        ["Вариант Option", "Физический эффект Physical effect", "Роль Role", "Решение Decision"],
        [
            ["4 x TDK T5838", "Компактный когерентный захват и AAD\nCompact coherent capture and AAD", "Базовый и контрольный канал\nBaseline and reference channel", "Сохранить\nKeep"],
            ["8-16 x Infineon IM73D122", "Ниже шум капсюля и больше апертура\nLower capsule noise and larger aperture", "Массив 80-2000 Hz\n80-2000 Hz array", "Прототип\nPrototype"],
            ["RODE NTG5 or Sennheiser MKH 416", "Интерференционная трубка\nInterference tube", "Направленный широкополосный эталон\nDirectional broadband reference", "Сравнить\nCompare"],
            ["Конический рупор 0.7 m\n0.7 m conical horn", "Осесимметричное преобразование давления\nAxisymmetric pressure transformation", "Эталон рупора 250-1500 Hz\nHorn reference 250-1500 Hz", "Изготовить один\nBuild one"],
            ["Пирамидальный рупор 0.8 x 0.5 m\nPyramidal horn", "Раздельные горизонтальная и вертикальная апертуры\nSeparate horizontal and vertical apertures", "Кандидат секторной головки\nSector-head candidate", "Изготовить один\nBuild one"],
            ["Отражатель 0.6 m\n0.6 m reflector", "Апертура с ростом усиления по частоте\nAperture gain increasing with frequency", "Опыт 500-3000 Hz\n500-3000 Hz experiment", "Изготовить один\nBuild one"],
            ["Сужающиеся четвертьволновые каналы\nTapered quarter-wave cells", "Узкополосное повышение давления\nNarrowband pressure gain", "Вспомогательные гармонические каналы\nAuxiliary harmonic channels", "Изготовить кассету\nBuild a bank"],
        ],
        [4.0, 5.4, 4.8, 3.9],
        7.0,
    )

    add_heading(document, "5 Первичная физика апертуры Aperture physics")
    add_bilingual_paragraph(
        document,
        "Длина волны составляет примерно 3,43 м на 100 Гц, 1,37 м на 250 Гц, 0,69 м на 500 Гц и 0,34 м на 1 кГц. Поэтому конструкция с раскрытием меньше метра не может создать узкую диаграмму в полосе 100-250 Гц. Для этой полосы основной физический инструмент представляет собой распределенный массив с базой 1,2-2 м, а не компактный рупор.",
        "Wavelength is approximately 3.43 m at 100 Hz, 1.37 m at 250 Hz, 0.69 m at 500 Hz and 0.34 m at 1 kHz. An aperture below one metre therefore cannot produce a narrow beam at 100-250 Hz. The primary physical instrument in that band is a distributed 1.2-2 m array, not a compact horn.",
    )
    add_bilingual_paragraph(
        document,
        "Секторная направленность начинает становиться практичной, когда размеры раскрыва сравнимы с длиной волны. Ниже этой области рупор можно использовать как согласующее устройство и защищенный акустический вход, но не следует приписывать ему узкую диаграмму без измерения.",
        "Sector directivity becomes practical when the aperture dimensions are comparable with wavelength. Below that region a horn may act as an impedance-transforming and protected acoustic inlet, but a narrow pattern must not be claimed without measurement.",
    )

    add_heading(document, "6 Конический и пирамидальный рупоры Conical and pyramidal horns")
    add_table(
        document,
        ["Параметр Parameter", "Конический Conical", "Пирамидальный Pyramidal", "Проверка Verification"],
        [
            ["Начальный раскрыв\nInitial mouth", "Диаметр около 0.7 m\nAbout 0.7 m diameter", "Около 0.8 x 0.5 m\nAbout 0.8 x 0.5 m", "Измерить площадь и фактический профиль\nMeasure area and actual flare"],
            ["Длина\nLength", "0.8-1.2 m", "0.8-1.2 m", "Одинаковая длина первого сравнения\nEqual first-comparison length"],
            ["Горло\nThroat", "30-60 mm", "30-60 mm", "Одинаковый капсюль и предусилитель\nSame capsule and preamplifier"],
            ["Диаграмма\nPattern", "Осесимметричная\nAxisymmetric", "Раздельная по горизонтали и вертикали\nSeparate horizontal and vertical", "12 азимутов и несколько углов места\n12 azimuths and several elevations"],
            ["Изготовление\nFabrication", "Криволинейная оболочка\nCurved shell", "Плоские панели\nFlat panels", "Масса, жесткость, дренаж, повторяемость\nMass, stiffness, drainage, repeatability"],
            ["Предварительный выбор\nInitial choice", "Эталон сравнения\nComparison reference", "Первый секторный кандидат\nFirst sector candidate", "Выбор только после одинакового теста\nSelect only after equal testing"],
        ],
        [3.7, 4.7, 5.2, 4.5],
        7.0,
    )
    add_bilingual_paragraph(
        document,
        "Пирамидальная форма выбрана первым секторным кандидатом, потому что горизонтальную и вертикальную апертуры можно менять независимо, а плоские панели проще изготовить, усилить и повторить. Конический вариант нужен как осесимметричный контроль. Преимущество любого профиля должно подтверждаться измеренной АЧХ, диаграммой, фазой, шумом потока и устойчивостью к дождю.",
        "The pyramidal shape is the first sector candidate because horizontal and vertical aperture can be adjusted independently, while flat panels are easier to build, stiffen and reproduce. The conical shape provides an axisymmetric control. Any profile advantage must be established by measured response, pattern, phase, flow noise and rain tolerance.",
    )

    add_heading(document, "7 Резонаторы и отражатели Resonators and reflectors")
    add_bilingual_paragraph(
        document,
        "Резонатор может повысить давление у капсюля в ограниченной полосе, но одновременно повышает добротность, время установления и чувствительность к геометрии, температуре, загрязнению и воде. Он не отделяет цель от внешнего шума той же частоты и направления. Поэтому кассета перекрывающихся сужающихся четвертьволновых каналов 100-800 Гц рассматривается как вспомогательный физический предфильтр гармоник.",
        "A resonator can increase pressure at the capsule in a limited band, but it also increases Q, settling time and sensitivity to geometry, temperature, contamination and water. It cannot separate a target from external noise at the same frequency and direction. An overlapping bank of tapered quarter-wave cells from 100-800 Hz is therefore treated as an auxiliary physical harmonic prefilter.",
    )
    add_bilingual_paragraph(
        document,
        "Отражатель диаметром 0,6 м дает полезную апертуру преимущественно в средней и верхней полосе. На низких частотах его размер мал относительно длины волны. Полевой вариант должен иметь дренаж, защиту от порывов, механическую развязку и индивидуальную калибровку положения капсюля.",
        "A 0.6 m reflector provides useful aperture mainly in the mid and upper band. At low frequencies it is small relative to wavelength. The field design needs drainage, gust protection, mechanical isolation and individual calibration of capsule position.",
    )

    add_heading(document, "8 Рекомендуемый прототип LR 1 Recommended LR 1 prototype")
    add_figure(
        document,
        LAYOUT_FIG,
        5.75,
        "Рисунок 2 Предварительная физическая компоновка LR-1",
        "Figure 2 Preliminary LR-1 physical layout",
    )
    steps = [
        ("Сохранить четыре штатных T5838 без изменения как AAD, резерв и контроль совместимости.", "Keep the four standard T5838 channels unchanged for AAD, backup and compatibility control."),
        ("Собрать низкочастотный массив с изменяемой базой 1,2-2 м для ориентировочной полосы 80-400 Гц.", "Build a low-frequency array with an adjustable 1.2-2 m baseline for an initial 80-400 Hz band."),
        ("Изготовить один конический и один пирамидальный рупор с одинаковым капсюлем, предусилителем и длиной.", "Build one conical and one pyramidal horn with the same capsule, preamplifier and length."),
        ("Сравнить рупоры с NTG5 или доступным NTG-2 с WS6, отражателем 0,6 м и кассетой резонаторов.", "Compare the horns with an NTG5 or available NTG-2 with WS6, a 0.6 m reflector and the resonator bank."),
        ("Использовать полноценную корзинную ветрозащиту, внешний мех, дренаж и развязку от вибрации мачты.", "Use a full basket windshield, external fur, drainage and mast-vibration isolation."),
        ("Первую серию записывать автономным синхронным восьмиканальным трактом 24 bit. PCB-MAIN менять только после выбора физической головки.", "Record the first series with an independent synchronous eight-channel 24-bit path. Change PCB-MAIN only after the physical head is selected."),
    ]
    for ru, en in steps:
        add_numbered_pair(document, ru, en)

    add_heading(document, "9 Программа физической проверки Physical verification program")
    add_table(
        document,
        ["Этап Stage", "Измерение Measurement", "Условия Conditions", "Результат Output"],
        [
            ["Лаборатория\nLaboratory", "АЧХ, собственный шум, диаграмма, фаза\nResponse, self-noise, pattern, phase", "80-3000 Hz, 12 азимутов, 3 уровня\n80-3000 Hz, 12 azimuths, 3 levels", "Калибровка каждого канала\nPer-channel calibration"],
            ["Ветер\nWind", "Шум 20-500 Hz, импульсы, клиппинг\n20-500 Hz noise, impulses, clipping", "0-2, 2-5 and 5-8 m/s", "Выбор защиты и развязки\nWindshield and isolation choice"],
            ["Дальность\nRange", "Сырой узкополосный SNR\nRaw narrowband SNR", "50, 100, 250, 500, 1000 and 2000 m", "SNR по частоте и дальности\nSNR by frequency and range"],
            ["Реальные цели\nReal targets", "Обнаружение, пропуски, курс, режим\nDetection, misses, aspect, state", "Все классы целей и фон\nAll target classes and background", "Pd и дальность по классу и погоде\nPd and range by class and weather"],
            ["Сравнение\nComparison", "T5838, shotgun, рупоры, отражатель, резонаторы\nT5838, shotgun, horns, reflector, resonators", "Одновременная запись на одной мачте\nSimultaneous recording on one mast", "Физический выигрыш без смены алгоритма\nPhysical gain without algorithm change"],
        ],
        [2.7, 5.2, 5.0, 5.2],
        7.0,
    )
    add_bilingual_paragraph(
        document,
        "Основная метрика первого этапа представляет собой разницу сырого SNR в одинаковых временных и частотных окнах. Один алгоритм применяется ко всем каналам без отдельной подстройки. Классификацию и пеленг сравнивают после физического измерения.",
        "The primary first-stage metric is the difference in raw SNR over identical time and frequency windows. One algorithm is applied to every channel without per-channel tuning. Classification and bearing performance are compared only after the physical measurement.",
    )

    add_heading(document, "10 Критерии решения Decision criteria")
    add_table(
        document,
        ["Критерий Criterion", "PASS", "HOLD", "Запись Record"],
        [
            ["Калибровка\nCalibration", "Полная по всем каналам\nComplete for all channels", "Нет уровня, фазы или геометрии\nMissing level, phase or geometry", "Файлы калибровки и SHA-256\nCalibration files and SHA-256"],
            ["Ветер\nWind", "Нет клиппинга, измерен SNR\nNo clipping, SNR measured", "Неопределенная защита или вибрация\nUnresolved shielding or vibration", "WAV, скорость и направление ветра\nWAV, wind speed and direction"],
            ["Дальность\nRange", "Pd и FAR утверждены для класса и режима\nPd and FAR approved for class and state", "Только единичное обнаружение\nSingle detection only", "Трасса, погода, фон и трек\nGeometry, weather, background and track"],
            ["Интеграция\nIntegration", "Питание, интерфейс и защита определены\nPower, interface and protection defined", "Требуется изменение серийной платы без ECO\nProduction-board change required without ECO", "Интерфейсный контрольный документ\nInterface control document"],
        ],
        [3.5, 5.2, 5.1, 4.3],
        7.0,
    )
    add_bilingual_paragraph(
        document,
        "Требование 2 км становится проверяемым только после фиксации минимального спектра источника на 1 м, скорости ветра, местности, температуры, влажности, высоты, сектора обзора, вероятности обнаружения и частоты ложных тревог. До этого дальность остается целью исследования, а не гарантированной характеристикой.",
        "The 2 km requirement becomes testable only after the minimum 1 m source spectrum, wind speed, terrain, temperature, humidity, altitude, field of view, probability of detection and false-alarm rate are fixed. Until then, range remains a research objective rather than a guaranteed characteristic.",
    )

    add_heading(document, "11 Следующий выпуск Next release")
    add_bilingual_paragraph(
        document,
        "Разрешить изготовление одного стенда LR-1. Состав стенда включает изменяемый низкочастотный массив, конический и пирамидальный рупоры, направленный эталонный микрофон, отражатель 0,6 м, кассету резонаторов, ветрозащиту и независимый синхронный рекордер. Выбор серийной физической головки выполняется после общей лабораторной и полевой серии.",
        "Authorize one LR-1 test stand. The stand includes an adjustable low-frequency array, conical and pyramidal horns, a directional reference microphone, a 0.6 m reflector, a resonator bank, wind protection and an independent synchronous recorder. The production physical head is selected after one common laboratory and field campaign.",
    )
    add_bilingual_paragraph(
        document,
        "Полные сравнительные таблицы, классы целей, расчет дальности и матрица испытаний находятся в книге EVT_PRE_20_ACOUSTIC_TRADE_STUDY_REV_C.xlsx. Машинночитаемые CSV сохраняются как приложения к книге.",
        "The full comparison tables, target classes, range calculation and test matrix are in EVT_PRE_20_ACOUSTIC_TRADE_STUDY_REV_C.xlsx. Machine-readable CSV files remain as workbook attachments.",
    )

    add_heading(document, "12 Источники Sources")
    add_sources(document)

    add_heading(document, "Приложение А Карта измерения цели Appendix A target measurement record")
    add_table(
        document,
        ["Поле Field", "Обязательные данные Required data", "Назначение Purpose"],
        [
            ["Идентификация\nIdentification", "Код, планер, силовая установка\nID, airframe, propulsion", "Связь спектра с конфигурацией\nTie spectrum to configuration"],
            ["Режим\nState", "Висение, набор, крейсер, маневр, удаление\nHover, climb, cruise, maneuver, retreat", "Раздельные дальности по состояниям\nSeparate range by state"],
            ["Геометрия\nGeometry", "Дальность, высота, азимут, угол места, курс, скорость\nRange, altitude, azimuth, elevation, course, speed", "Пересчет и ракурс\nNormalization and aspect"],
            ["Источник\nSource", "Обороты, винты или роторы, двигатель и выхлоп\nSpeed, propellers or rotors, engine and exhaust", "Связь линий с физикой\nTie spectral lines to physics"],
            ["Среда\nEnvironment", "Температура, влажность, давление, ветер, поверхность\nTemperature, humidity, pressure, wind, ground", "Потери и фон\nLoss and background"],
            ["Запись\nRecording", "Калиброванный WAV, чувствительность, шум, синхронизация\nCalibrated WAV, sensitivity, noise, synchronization", "Повторяемое сравнение\nRepeatable comparison"],
            ["Результат\nResult", "SNR, диаграмма, обнаружение, пропуски, FAR\nSNR, pattern, detection, misses, FAR", "Решение по физической головке\nPhysical-head decision"],
        ],
        [3.8, 9.0, 5.4],
        7.1,
    )

    add_heading(document, "Приложение Б История изменений Appendix B revision history")
    add_table(
        document,
        ["Ревизия Revision", "Дата Date", "Изменение Change", "Статус Status"],
        [
            ["Rev A", "27.10.2026", "Первичное сравнение физической акустики\nInitial physical-acoustics comparison", "Сохранено\nRetained"],
            ["Rev B", "27.10.2026", "Добавлены рупоры и классы целей\nAdded horn profiles and target classes", "Сохранено\nRetained"],
            ["Rev C", "27.10.2026", "Переиздано в контролируемом двуязычном формате EVT, добавлены XLSX и CSV приложения\nReissued in the controlled bilingual EVT format with XLSX and CSV attachments", "Текущая\nCurrent"],
        ],
        [3.0, 3.0, 9.8, 2.4],
        7.2,
    )

    document.core_properties.title = "Альтернативный акустический тракт Дионея для дальности от 2 км"
    document.core_properties.subject = "Dioneya EVT-PRE-20 LR-1 acoustic trade study Rev C"
    document.core_properties.comments = "Controlled bilingual engineering evaluation. EVT production baseline unchanged."
    document.save(DOCX)


def build_markdown() -> None:
    text = f"""# Альтернативный акустический тракт Дионея для дальности от 2 км

## Dioneya alternative long range acoustic front end

**Документ / Document:** {STEM}

**Ревизия / Revision:** Rev C

**Дата / Date:** 27.10.2026

**Статус / Status:** ENGINEERING EVALUATION; EVT PRODUCTION BASELINE UNCHANGED

## Решение Decision

RU: Штатная головка из четырех TDK T5838 сохраняется в EVT-PRE-20 как малопотребляющий дежурный тракт, AAD и контрольный канал. Для проверки дальности 2 км нужен отдельный прототип LR-1 с большой апертурой, низким собственным шумом, секторной направленностью, ветрозащитой и калиброванной записью.

EN: The standard four-channel TDK T5838 head remains in EVT-PRE-20 as the low-power watch channel, AAD and reference path. Verification at 2 km requires a separate LR-1 prototype with a large aperture, low self-noise, sector directivity, wind protection and calibrated recording.

RU: LR-1 является параллельным экспериментальным модулем и не изменяет выпущенную производственную конфигурацию EVT-PRE-20.

EN: LR-1 is a parallel experimental module and does not change the released EVT-PRE-20 production configuration.

## Ключевая физика Key physical limits

RU: На 2 км сферическое расхождение дает 66,0 дБ потерь. Длина волны составляет примерно 3,43 м на 100 Гц, 1,37 м на 250 Гц, 0,69 м на 500 Гц и 0,34 м на 1 кГц. Рупор с раскрытием меньше метра не создает узкую диаграмму в полосе 100-250 Гц. Для этой полосы нужен распределенный массив с базой 1,2-2 м.

EN: At 2 km, spherical spreading produces 66.0 dB of loss. Wavelength is approximately 3.43 m at 100 Hz, 1.37 m at 250 Hz, 0.69 m at 500 Hz and 0.34 m at 1 kHz. A sub-metre horn cannot produce a narrow beam at 100-250 Hz. That band requires a distributed 1.2-2 m array.

## Следующий шаг Next step

RU: Изготовить один стенд LR-1 с изменяемым низкочастотным массивом, коническим и пирамидальным рупорами, направленным эталонным микрофоном, отражателем, кассетой резонаторов, ветрозащитой и независимым синхронным рекордером.

EN: Build one LR-1 stand with an adjustable low-frequency array, conical and pyramidal horns, a directional reference microphone, reflector, resonator bank, wind protection and an independent synchronous recorder.
"""
    MD.write_text(text, encoding="utf-8")


build_figures()
build_document()
build_markdown()
print(DOCX)
print(MD)
