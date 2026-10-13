from __future__ import annotations

import argparse
import csv
import io
import json
import re
import zipfile
from collections import Counter
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from openpyxl import load_workbook
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_RELEASE = ROOT / "releases" / "evt-pre-20" / "2026-10-13-rev-d6"
ERROR_TOKENS = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!", "#SPILL!", "#CALC!")


def iter_delivery_files(release: Path):
    for folder in (
        "Основной_пакет",
        "Заказ_плат_и_монтаж",
        "Заказ_жгутов_и_кабелей",
        "Сборка_и_корпуса",
        "Программа_EVT_и_Мухоед",
        "Ревью_PCB_MIC_отдельно",
    ):
        root = release / folder
        for path in sorted(root.rglob("*")):
            if path.is_file() and path.suffix.lower() != ".zip":
                yield path


def in_table(paragraph) -> bool:
    return paragraph._p.getparent().tag.endswith("}tc")


def effective_style_value(style, attribute):
    current = style
    while current is not None:
        value = getattr(current.font, attribute)
        if value is not None:
            return value
        current = current.base_style
    return None


def effective_run_value(run, paragraph, attribute):
    direct = getattr(run.font, attribute)
    if direct is not None:
        return direct
    return effective_style_value(paragraph.style, attribute)


def inspect_docx(path: Path) -> dict:
    document = Document(path)
    all_paragraphs = list(document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                all_paragraphs.extend(cell.paragraphs)

    em_dash = 0
    explicit_wrong_font = 0
    explicit_wrong_size = 0
    body_alignment_flags = 0
    body_spacing_flags = 0
    heading_not_bold = 0
    text_runs = 0
    for paragraph in all_paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        em_dash += text.count("—")
        style_name = paragraph.style.name.lower() if paragraph.style is not None else ""
        heading = style_name.startswith("heading") or style_name.startswith("заголов")
        table_text = in_table(paragraph)
        body = not heading and not table_text and style_name not in {"title", "subtitle", "название"}
        if body and paragraph.alignment not in (None, WD_ALIGN_PARAGRAPH.JUSTIFY):
            body_alignment_flags += 1
        if body:
            spacing = paragraph.paragraph_format.line_spacing
            if spacing not in (None, 1.0):
                body_spacing_flags += 1
        paragraph_has_bold = False
        for run in paragraph.runs:
            if not run.text.strip():
                continue
            text_runs += 1
            name = effective_run_value(run, paragraph, "name")
            size = effective_run_value(run, paragraph, "size")
            bold = effective_run_value(run, paragraph, "bold")
            paragraph_has_bold = paragraph_has_bold or bool(bold)
            if run.font.name is not None and name and name.lower() != "times new roman":
                explicit_wrong_font += 1
            expected_pt = 12.0 if table_text else 14.0
            if run.font.size is not None and size is not None and abs(size.pt - expected_pt) > 0.1:
                explicit_wrong_size += 1
        if heading and not paragraph_has_bold:
            heading_not_bold += 1

    normal = document.styles["Normal"]
    normal_font = effective_style_value(normal, "name")
    normal_size = effective_style_value(normal, "size")
    return {
        "paragraphs": len(all_paragraphs),
        "tables": len(document.tables),
        "text_runs": text_runs,
        "normal_font": normal_font,
        "normal_size_pt": normal_size.pt if normal_size is not None else None,
        "em_dash_count": em_dash,
        "explicit_wrong_font_runs": explicit_wrong_font,
        "explicit_wrong_size_runs": explicit_wrong_size,
        "body_alignment_flags": body_alignment_flags,
        "body_spacing_flags": body_spacing_flags,
        "heading_not_bold_flags": heading_not_bold,
    }


def inspect_xlsx(path: Path) -> dict:
    workbook = load_workbook(path, read_only=True, data_only=False)
    formula_errors = []
    nonempty = 0
    keyword_hits: dict[str, list[str]] = {key: [] for key in ("аккумуля", "солнеч", "антенн", "rf", "кабел")}
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                value = cell.value
                if value is None:
                    continue
                nonempty += 1
                text = str(value)
                if any(token in text for token in ERROR_TOKENS):
                    formula_errors.append(f"{sheet.title}!{cell.coordinate}={text}")
                lower = text.lower()
                for keyword in keyword_hits:
                    if keyword in lower and len(keyword_hits[keyword]) < 20:
                        keyword_hits[keyword].append(f"{sheet.title}!{cell.coordinate}={text[:160]}")
    workbook.close()
    return {
        "sheets": list(workbook.sheetnames),
        "nonempty_cells": nonempty,
        "formula_errors": formula_errors,
        "keyword_hits": {key: value for key, value in keyword_hits.items() if value},
    }


def inspect_pdf(path: Path) -> dict:
    reader = PdfReader(path)
    if reader.is_encrypted:
        raise ValueError(f"encrypted PDF: {path}")
    texts = [(page.extract_text() or "") for page in reader.pages]
    text = "\n".join(texts)
    cyrillic = len(re.findall(r"[А-Яа-яЁё]", text))
    latin = len(re.findall(r"[A-Za-z]", text))
    bilingual_expected = "RU_EN" in path.name
    bilingual_pass = not bilingual_expected or (cyrillic >= 20 and latin >= 20)
    return {
        "pages": len(reader.pages),
        "text_characters": len(text),
        "cyrillic_characters": cyrillic,
        "latin_characters": latin,
        "bilingual_expected": bilingual_expected,
        "bilingual_pass": bilingual_pass,
    }


def inspect_csv(path: Path) -> dict:
    data = path.read_bytes()
    text = None
    encoding_used = None
    for encoding in ("utf-8-sig", "utf-8", "cp1251"):
        try:
            text = data.decode(encoding)
            encoding_used = encoding
            break
        except UnicodeDecodeError:
            pass
    if text is None:
        raise ValueError(f"CSV decode failure: {path}")
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ","
    rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    return {
        "encoding": encoding_used,
        "delimiter": delimiter,
        "rows": len(rows),
        "max_columns": max((len(row) for row in rows), default=0),
    }


def validate_zip(path: Path) -> dict:
    with zipfile.ZipFile(path) as archive:
        bad = archive.testzip()
        files = [item for item in archive.infolist() if not item.is_dir()]
        names = [item.filename for item in files]
        return {
            "entries": len(files),
            "crc_pass": bad is None,
            "duplicate_paths": len(names) - len(set(names)),
        }


def main() -> None:
    parser = argparse.ArgumentParser(description="Deep audit of the EVT-PRE-20 Rev D6 delivery")
    parser.add_argument("--release", type=Path, default=DEFAULT_RELEASE)
    parser.add_argument("--out", type=Path)
    args = parser.parse_args()
    release = args.release.resolve()
    output = args.out or release / "ПРОТОКОЛ_ПОЛНОЙ_ПРОВЕРКИ_REV_D6.json"

    documents = {}
    workbooks = {}
    pdfs = {}
    csvs = {}
    errors = []
    extensions: Counter[str] = Counter()
    for path in iter_delivery_files(release):
        relative = path.relative_to(release).as_posix()
        suffix = path.suffix.lower()
        extensions[suffix or "<none>"] += 1
        try:
            if suffix == ".docx":
                documents[relative] = inspect_docx(path)
            elif suffix == ".xlsx":
                workbooks[relative] = inspect_xlsx(path)
            elif suffix == ".pdf":
                pdfs[relative] = inspect_pdf(path)
            elif suffix == ".csv":
                csvs[relative] = inspect_csv(path)
        except Exception as exc:
            errors.append(f"{relative}: {exc}")

    archives = {}
    for path in sorted(release.glob("*.zip")):
        try:
            archives[path.name] = validate_zip(path)
        except Exception as exc:
            errors.append(f"{path.name}: {exc}")

    docx_summary = {
        "files": len(documents),
        "open_pass": not any(error.lower().endswith(".docx") for error in errors),
        "em_dash_count": sum(item["em_dash_count"] for item in documents.values()),
        "explicit_wrong_font_runs": sum(item["explicit_wrong_font_runs"] for item in documents.values()),
        "explicit_wrong_size_runs": sum(item["explicit_wrong_size_runs"] for item in documents.values()),
        "body_alignment_flags": sum(item["body_alignment_flags"] for item in documents.values()),
        "body_spacing_flags": sum(item["body_spacing_flags"] for item in documents.values()),
        "heading_not_bold_flags": sum(item["heading_not_bold_flags"] for item in documents.values()),
    }
    xlsx_summary = {
        "files": len(workbooks),
        "formula_errors": sum(len(item["formula_errors"]) for item in workbooks.values()),
    }
    pdf_summary = {
        "files": len(pdfs),
        "pages": sum(item["pages"] for item in pdfs.values()),
        "bilingual_failures": [name for name, item in pdfs.items() if not item["bilingual_pass"]],
    }

    procurement = next(
        (item for name, item in workbooks.items() if "Закупочная ведомость_BOM" in name),
        None,
    )
    procurement_terms = procurement["keyword_hits"] if procurement else {}
    required_terms = {"аккумуля", "солнеч", "антенн", "кабел"}
    missing_terms = sorted(required_terms - set(procurement_terms))

    critical = list(errors)
    if xlsx_summary["formula_errors"]:
        critical.append("spreadsheet formula errors detected")
    if pdf_summary["bilingual_failures"]:
        critical.append("RU_EN PDF without both language character sets")
    if missing_terms:
        critical.append(f"procurement workbook lacks terms: {missing_terms}")
    for name, item in archives.items():
        if not item["crc_pass"] or item["duplicate_paths"]:
            critical.append(f"archive integrity failure: {name}")

    report = {
        "schema": "dioneya-evt-pre-20-deep-package-audit-v1",
        "release": "EVT-PRE-20 Rev D6",
        "date": "2026-10-13",
        "result": "PASS" if not critical else "FAIL",
        "scope": "unpacked delivery folders and all release archives",
        "files_by_extension": dict(sorted(extensions.items())),
        "docx_summary": docx_summary,
        "xlsx_summary": xlsx_summary,
        "pdf_summary": pdf_summary,
        "procurement_required_terms": {
            "result": "PASS" if not missing_terms else "FAIL",
            "hits": procurement_terms,
            "missing": missing_terms,
        },
        "archives": archives,
        "errors": errors,
        "critical": critical,
        "documents": documents,
        "workbooks": workbooks,
        "pdfs": pdfs,
        "csvs": csvs,
    }
    output.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({key: report[key] for key in (
        "result", "files_by_extension", "docx_summary", "xlsx_summary", "pdf_summary",
        "procurement_required_terms", "archives", "critical"
    )}, ensure_ascii=False, indent=2))
    if critical:
        raise SystemExit(2)


if __name__ == "__main__":
    main()
