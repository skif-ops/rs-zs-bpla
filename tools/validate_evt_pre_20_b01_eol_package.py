#!/usr/bin/env python3
"""Validate the MFG-004 and DIO-EVT-B01 physical qualification package."""

from __future__ import annotations

import csv
import hashlib
import json
from pathlib import Path
import re
import sys
import zipfile

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
MFG = ROOT / "manufacturing"
DOCS = ROOT / "docs"
PROGRAM = MFG / "MFG_004_DIO_EVT_B01_EOL_PROGRAM_REV_A.json"
CONTACTS = MFG / "MFG_004_EOL_FIXTURE_CONTACT_MAP_REV_A.csv"
REGISTER_CSV = MFG / "MFG_004_DIO_EVT_B01_EOL_MEASUREMENT_REGISTER_REV_A.csv"
REGISTER_XLSX = MFG / "MFG_004_DIO_EVT_B01_EOL_MEASUREMENT_REGISTER_REV_A.xlsx"
DOC_MD = DOCS / "EVT_PRE_20_DIO_EVT_B01_PHYSICAL_EOL_PROGRAM_REV_A.md"
DOC_DOCX = DOCS / "EVT_PRE_20_DIO_EVT_B01_PHYSICAL_EOL_PROGRAM_REV_A.docx"
FIXTURE_PDF = MFG / "MFG_004_EOL_FIXTURE_DRILL_TEMPLATE_REV_A.pdf"
FIXTURE_DXF = MFG / "MFG_004_EOL_FIXTURE_DRILL_TEMPLATE_REV_A.dxf"
ARCHIVE = ROOT / "outputs" / "EVT_PRE_20_DIO_EVT_B01_И_EOL_ОСНАСТКА_20261009_REV_A.zip"
SHA256 = re.compile(r"^[0-9a-f]{64}$")


def require(condition: bool, message: str) -> None:
    if not condition:
        raise ValueError(message)


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def validate_sources() -> None:
    program = json.loads(PROGRAM.read_text(encoding="utf-8"))
    require(program["schema"] == 1, "unexpected program schema")
    require(program["program_id"] == "MFG_004_DIO_EVT_B01_EOL_PROGRAM_REV_A", "unexpected program ID")
    require(program["unit"] == {"serial": "DIO-EVT-B01", "station_id": 901, "tenant": "bench"}, "unexpected B01 identity")
    require(program["release_id"] == "EVT_PRE_20_BENCH_RELEASE_2026100601", "unexpected release ID")
    require(program["decision_policy"]["firmware_lock_authorization"] == "DENIED_UNTIL_SEPARATE_PRODUCTION_RELEASE", "lock interlock was relaxed")
    contacts = read_csv(CONTACTS)
    require(len(contacts) == 31, f"expected 31 fixture contacts, found {len(contacts)}")
    expected_counts = {"TP_EOL": 13, "TP_MCU_SWD": 5, "TP_BLE_SWD": 4, "TP_CELL_USB": 4, "TP_CELL_DBG": 5}
    actual_counts = {group: sum(row["Fixture_Group"] == group for row in contacts) for group in expected_counts}
    require(actual_counts == expected_counts, f"contact groups mismatch: {actual_counts}")
    require(len(program["contacts"]) == len(contacts), "JSON contact map count differs from CSV")
    for item, csv_item in zip(program["contacts"], contacts, strict=True):
        for key in ("Fixture_Group", "Contact", "Board_View_X_mm", "Board_View_Y_mm", "Bottom_Fixture_View_X_mm", "Bottom_Fixture_View_Y_mm", "Pin_Name", "Net", "Safety_Class"):
            require(str(item[key]) == str(csv_item[key]), f"contact mismatch {key}: {item[key]} != {csv_item[key]}")
        expected_fixture_x = 110.0 - float(csv_item["Board_View_X_mm"])
        require(abs(expected_fixture_x - float(csv_item["Bottom_Fixture_View_X_mm"])) < 0.001, "bottom fixture mirror transform mismatch")
    steps = program["steps"]
    require(len(steps) >= 35, "physical EOL sequence is too short")
    ids = [item["test_id"] for item in steps]
    require(len(ids) == len(set(ids)), "duplicate program step ID")
    require(all(item["mandatory"] is True for item in steps), "mandatory step was relaxed")
    require(all(item["initial_status"] == "NOT_RUN" for item in steps), "source program contains a non-NOT_RUN result")
    require(any(item["limit_state"] == "OPEN_B01_FREEZE" for item in steps), "open physical limits were hidden")
    require(any(item["test_id"] == "PWR-3V8-01" and item["lower_limit"] == 3.3 for item in steps), "approved BG95 minimum was lost")
    require(any(item["test_id"] == "MIC-CLK-01" and item["lower_limit"] == 50000 for item in steps), "approved MIC clock minimum was lost")
    require(len(program["msa"]) >= 8 and all(item["status"] == "NOT_RUN" for item in program["msa"]), "MSA matrix is incomplete or pre-approved")
    register = read_csv(REGISTER_CSV)
    require(len(register) == len(steps), "CSV register row count mismatch")
    require(all(item["status"] == "NOT_RUN" for item in register), "CSV register contains a result")
    md = DOC_MD.read_text(encoding="utf-8")
    for character in ("\u2014", "\u2013", "\u2011"):
        require(character not in md, f"forbidden dash U+{ord(character):04X} in document source")
    for token in ("TP_EOL.2", "TP_MCU_SWD.4", "TP_CELL_DBG.4", "TP_CELL_USB.1", "J_MIC1-J_MIC4", "OPEN_B01_FREEZE"):
        require(token in md, f"instruction is missing required token {token}")


def validate_xlsx() -> None:
    workbook = load_workbook(REGISTER_XLSX, data_only=False)
    require(workbook.sheetnames == ["Summary", "Contacts", "Measurements", "MSA", "Evidence"], "unexpected workbook sheets")
    measurements = workbook["Measurements"]
    rows = 0
    for row in range(5, measurements.max_row + 1):
        if measurements.cell(row, 2).value:
            rows += 1
            require(measurements.cell(row, 24).value == "NOT_RUN", f"row {row} is not NOT_RUN")
            require(str(measurements.cell(row, 27).value).startswith("=IF("), f"row {row} gate formula is missing")
    program = json.loads(PROGRAM.read_text(encoding="utf-8"))
    require(rows == len(program["steps"]), "XLSX measurement row count mismatch")
    summary = workbook["Summary"]
    require(str(summary["B26"].value).startswith("=IF("), "overall gate formula is missing")
    cached = load_workbook(REGISTER_XLSX, data_only=True)
    cached_summary = cached["Summary"]
    open_limits = sum(
        "OPEN" in str(item["limit_state"]) or str(item["limit_state"]).startswith("PARTIAL")
        for item in program["steps"]
    )
    require(cached_summary["B20"].value == open_limits, "cached OPEN limit count is stale")
    require(cached_summary["B26"].value == "NOT_RUN", "initial overall state is not NOT_RUN")


def validate_fixture_outputs() -> None:
    require(FIXTURE_PDF.is_file() and FIXTURE_PDF.stat().st_size > 10_000, "fixture PDF is missing or empty")
    require(FIXTURE_DXF.is_file() and FIXTURE_DXF.stat().st_size > 1_000, "fixture DXF is missing or empty")
    dxf = FIXTURE_DXF.read_text(encoding="ascii")
    require(dxf.count("\nCIRCLE\n") >= 35, "fixture DXF does not contain 31 pogo holes and 4 mounting holes")


def validate_archive() -> None:
    if not ARCHIVE.exists():
        return
    with zipfile.ZipFile(ARCHIVE) as archive:
        names = set(archive.namelist())
        required_suffixes = {
            "EVT_PRE_20_DIO_EVT_B01_PHYSICAL_EOL_PROGRAM_REV_A.docx",
            "MFG_004_DIO_EVT_B01_EOL_MEASUREMENT_REGISTER_REV_A.xlsx",
            "MFG_004_DIO_EVT_B01_EOL_MEASUREMENT_REGISTER_REV_A.csv",
            "MFG_004_EOL_FIXTURE_CONTACT_MAP_REV_A.csv",
            "MFG_004_EOL_FIXTURE_DRILL_TEMPLATE_REV_A.pdf",
            "MFG_004_EOL_FIXTURE_DRILL_TEMPLATE_REV_A.dxf",
            "MFG_004_DIO_EVT_B01_EOL_PROGRAM_REV_A.json",
            "EVT_PRE_20_B01_HARDWARE_QUALIFICATION_REV_A.xlsx",
            "EVT_PRE_20_BENCH_RELEASE_2026100601.zip",
            "SHA256SUMS.txt",
        }
        for suffix in required_suffixes:
            require(any(name.endswith(suffix) for name in names), f"archive is missing {suffix}")
        checksum_name = next(name for name in names if name.endswith("SHA256SUMS.txt"))
        checksum_lines = archive.read(checksum_name).decode("utf-8").splitlines()
        for line in checksum_lines:
            digest, name = line.split("  ", 1)
            require(SHA256.fullmatch(digest) is not None, f"bad archive hash for {name}")
            require(name in names, f"hash references missing archive member {name}")
            require(hashlib.sha256(archive.read(name)).hexdigest() == digest, f"archive hash mismatch for {name}")


def main() -> int:
    validate_sources()
    require(DOC_DOCX.exists(), "DOCX was not built")
    require(REGISTER_XLSX.exists(), "XLSX was not built")
    validate_xlsx()
    validate_fixture_outputs()
    validate_archive()
    print("DIO-EVT-B01 physical program and MFG-004 EOL package: PASS")
    print("physical results: NOT_RUN")
    print("firmware locking authorized: NO")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except (KeyError, OSError, ValueError, zipfile.BadZipFile) as exc:
        print(f"DIO-EVT-B01 EOL package: FAIL: {exc}")
        raise SystemExit(1) from exc
