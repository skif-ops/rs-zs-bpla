from __future__ import annotations

import math
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


REPO_ROOT = Path(__file__).resolve().parent.parent
DOCS = REPO_ROOT / "docs"
FIGURES = DOCS / "figures"
DOCS.mkdir(parents=True, exist_ok=True)
FIGURES.mkdir(parents=True, exist_ok=True)

STEM = "EVT_PRE_20_LONG_RANGE_ACOUSTIC_HEAD_TRADE_STUDY_REV_A"
DOCX = DOCS / f"{STEM}.docx"
MD = DOCS / f"{STEM}.md"
LOSS_FIG = FIGURES / "acoustic_path_loss_2km.png"
LAYOUT_FIG = FIGURES / "recommended_acoustic_head_layout.png"


def iso9613_absorption_db_per_m(freq_hz: float, temp_c: float = 15.0, rh_percent: float = 70.0) -> float:
    temperature = 273.15 + temp_c
    reference_temperature = 293.15
    triple_point = 273.16
    pressure_ratio = 1.0
    saturation_ratio = 10 ** (-6.8346 * (triple_point / temperature) ** 1.261 + 4.6151)
    molar_humidity_percent = rh_percent * saturation_ratio / pressure_ratio
    oxygen_relaxation = pressure_ratio * (
        24.0
        + 4.04e4
        * molar_humidity_percent
        * (0.02 + molar_humidity_percent)
        / (0.391 + molar_humidity_percent)
    )
    nitrogen_relaxation = (
        pressure_ratio
        * (temperature / reference_temperature) ** -0.5
        * (
            9.0
            + 280.0
            * molar_humidity_percent
            * math.exp(-4.17 * ((temperature / reference_temperature) ** (-1.0 / 3.0) - 1.0))
        )
    )
    frequency_sq = freq_hz**2
    classical = 1.84e-11 * pressure_ratio**-1 * (temperature / reference_temperature) ** 0.5
    molecular = (temperature / reference_temperature) ** -2.5 * (
        0.01275
        * math.exp(-2239.1 / temperature)
        / (oxygen_relaxation + frequency_sq / oxygen_relaxation)
        + 0.1068
        * math.exp(-3352.0 / temperature)
        / (nitrogen_relaxation + frequency_sq / nitrogen_relaxation)
    )
    return 8.686 * frequency_sq * (classical + molecular)


def path_loss_db(freq_hz: float, distance_m: float = 2000.0) -> float:
    return 20.0 * math.log10(distance_m) + iso9613_absorption_db_per_m(freq_hz) * distance_m


def ideal_dish_gain_db(freq_hz: float, diameter_m: float = 0.6, efficiency: float = 0.5) -> float:
    wavelength = 343.0 / freq_hz
    linear = efficiency * (math.pi * diameter_m / wavelength) ** 2
    return max(0.0, 10.0 * math.log10(linear))


FREQUENCIES = (100, 250, 500, 1000, 2000, 4000)
SOURCE_LEVELS = (70, 85, 100)


def build_figures() -> None:
    font_path = Path("C:/Windows/Fonts/arial.ttf")
    bold_path = Path("C:/Windows/Fonts/arialbd.ttf")
    font = ImageFont.truetype(str(font_path), 24)
    small = ImageFont.truetype(str(font_path), 20)
    tiny = ImageFont.truetype(str(font_path), 17)
    bold = ImageFont.truetype(str(bold_path), 28)

    image = Image.new("RGB", (1500, 820), "white")
    draw = ImageDraw.Draw(image)
    left, top, right, bottom = 135, 95, 1440, 690
    x_min, x_max = math.log10(80), math.log10(5000)
    y_min, y_max = 60.0, 145.0

    def point(freq: float, value: float) -> tuple[int, int]:
        x = left + (math.log10(freq) - x_min) / (x_max - x_min) * (right - left)
        y = bottom - (value - y_min) / (y_max - y_min) * (bottom - top)
        return int(x), int(y)

    for value in range(60, 146, 10):
        _, y = point(80, value)
        draw.line((left, y, right, y), fill="#D7DEE7", width=1)
        draw.text((55, y - 12), str(value), font=small, fill="#1F2937")
    for freq in (80, 100, 200, 500, 1000, 2000, 5000):
        x, _ = point(freq, 60)
        draw.line((x, top, x, bottom), fill="#E5E7EB", width=1)
        label = str(freq)
        box = draw.textbbox((0, 0), label, font=small)
        draw.text((x - (box[2] - box[0]) / 2, bottom + 12), label, font=small, fill="#1F2937")
    draw.line((left, top, left, bottom), fill="#111827", width=2)
    draw.line((left, bottom, right, bottom), fill="#111827", width=2)

    series = [
        ("Только сферическое расхождение", "#6B7280", [20.0 * math.log10(2000.0)] * 401),
        ("Расхождение и поглощение воздуха", "#123B63", []),
        ("Идеальная оценка с отражателем 0,6 м", "#B45309", []),
    ]
    frequencies = [10 ** (x_min + index * (x_max - x_min) / 400.0) for index in range(401)]
    series[1][2].extend(path_loss_db(freq) for freq in frequencies)
    series[2][2].extend(path_loss_db(freq) - ideal_dish_gain_db(freq) for freq in frequencies)
    for _, color, values in series:
        coordinates = [point(freq, value) for freq, value in zip(frequencies, values)]
        draw.line(coordinates, fill=color, width=5)

    draw.text((left, 30), "Расчёт потерь при 15 °C и относительной влажности 70 %", font=bold, fill="#111827")
    draw.text(((left + right) // 2 - 75, 752), "Частота, Гц", font=font, fill="#111827")
    draw.text((12, 335), "Потери, дБ", font=font, fill="#111827")
    legend_y = 118
    for label, color, _ in series:
        draw.line((left + 25, legend_y + 10, left + 95, legend_y + 10), fill=color, width=6)
        draw.text((left + 110, legend_y - 3), label, font=tiny, fill="#111827")
        legend_y += 32
    image.save(LOSS_FIG)

    image = Image.new("RGB", (1200, 1200), "white")
    draw = ImageDraw.Draw(image)
    center_x, center_y = 600, 560
    scale = 500
    radius_lower = 0.60 * scale
    radius_upper = 0.45 * scale
    draw.ellipse(
        (center_x - radius_lower, center_y - radius_lower, center_x + radius_lower, center_y + radius_lower),
        outline="#9CA3AF",
        width=3,
    )

    def arrow(angle_deg: float, radius: float, length: float, color: str) -> None:
        angle = math.radians(angle_deg)
        x = center_x + radius * math.cos(angle)
        y = center_y - radius * math.sin(angle)
        end_x = x + length * math.cos(angle)
        end_y = y - length * math.sin(angle)
        draw.line((x, y, end_x, end_y), fill=color, width=12)
        head = 25
        for offset in (-0.45, 0.45):
            head_angle = angle + math.pi + offset
            draw.line((end_x, end_y, end_x + head * math.cos(head_angle), end_y - head * math.sin(head_angle)), fill=color, width=10)
        draw.ellipse((x - 16, y - 16, x + 16, y + 16), fill=color)

    for angle in (0, 90, 180, 270):
        arrow(angle, radius_lower, 145, "#123B63")
    for angle in (45, 135, 225, 315):
        arrow(angle, radius_upper, 125, "#B45309")
    draw.rectangle((center_x - 18, center_y - 18, center_x + 18, center_y + 18), fill="#111827")
    draw.text((center_x + 28, center_y - 15), "мачта", font=font, fill="#111827")
    draw.text((90, 45), "Предлагаемая восьмисекторная акустическая головка", font=bold, fill="#111827")
    draw.text((90, 85), "вид сверху", font=font, fill="#111827")
    draw.line((95, 1040, 165, 1040), fill="#123B63", width=12)
    draw.text((185, 1025), "нижний ярус, наклон 10°", font=small, fill="#111827")
    draw.line((630, 1040, 700, 1040), fill="#B45309", width=12)
    draw.text((720, 1025), "верхний ярус, наклон 40–45°", font=small, fill="#111827")
    image.save(LAYOUT_FIG)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 90, start: int = 100, bottom: int = 90, end: int = 100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_keep_with_next(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    keep_next = OxmlElement("w:keepNext")
    p_pr.append(keep_next)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_1)
    run._r.append(instr_text)
    run._r.append(fld_char_2)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths_cm: list[float], font_size: float = 8.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, text in enumerate(headers):
        cell = header.cells[index]
        cell.width = Cm(widths_cm[index])
        set_cell_shading(cell, "17365D")
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(text)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(font_size)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cell = cells[index]
            cell.width = Cm(widths_cm[index])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            if row_index % 2:
                set_cell_shading(cell, "EDF3F8")
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            if index == 0 and len(headers) > 2:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(str(value))
            run.font.size = Pt(font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_bullet(doc: Document, text: str, level: int = 0) -> None:
    style = "List Bullet" if level == 0 else "List Bullet 2"
    paragraph = doc.add_paragraph(style=style)
    paragraph.add_run(text)


def add_number(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Number")
    paragraph.add_run(text)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    section.header_distance = Cm(0.7)
    section.footer_distance = Cm(0.7)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08
    for name, size, before, after in (
        ("Title", 20, 0, 12),
        ("Heading 1", 14, 12, 6),
        ("Heading 2", 11.5, 9, 4),
        ("Heading 3", 10.2, 7, 3),
    ):
        style = styles[name]
        style.font.name = "Aptos Display" if name != "Normal" else "Aptos"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = name != "Title" or True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
    for style_name in ("List Bullet", "List Bullet 2", "List Number"):
        styles[style_name].font.name = "Aptos"
        styles[style_name].font.size = Pt(9.5)
        styles[style_name].font.color.rgb = RGBColor(0, 0, 0)

    header = section.header.paragraphs[0]
    header.text = "Дионея EVT-PRE-20    Альтернативный акустический тракт    Rev A"
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor(90, 90, 90)
    add_page_number(section.footer.paragraphs[0])


def paragraph(doc: Document, text: str, bold_lead: str | None = None) -> None:
    item = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        first, rest = text.split(bold_lead, 1)
        item.add_run(first + bold_lead).bold = True
        item.add_run(rest)
    else:
        item.add_run(text)


def build_docx() -> None:
    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)
    title_run = title.add_run("Альтернативный акустический тракт Дионея для дальности от 2 км")
    title_run.font.name = "Aptos Display"
    title_run.font.size = Pt(20)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(0, 0, 0)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Инженерное сравнение первичной акустики и план физического прототипа")
    run.bold = True
    run.font.size = Pt(11)

    add_table(
        doc,
        ["Поле", "Значение"],
        [
            ["Документ", STEM],
            ["Ревизия и дата", "Rev A, 27.10.2026"],
            ["Исходная конфигурация", "4 × TDK T5838, геометрия 3+1, максимальная база 0,165 м"],
            ["Статус", "Инженерное исследование. Производственная конфигурация не изменена"],
        ],
        [4.3, 12.2],
        8.8,
    )

    intro = doc.add_paragraph()
    intro.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    intro.add_run("Главный вывод. ").bold = True
    intro.add_run(
        "Текущая четырёхмикрофонная головка T5838 подходит для малопотребляющего дежурного режима, "
        "распознавания в достижимом диапазоне и оценки направления. Она не даёт физического основания обещать "
        "распознавание малых электрических БПЛА на 2 км. Для такой дальности нужен отдельный крупноапертурный "
        "тракт с низким собственным шумом, выраженной направленностью и полноценной ветрозащитой. "
        "Существующий узел следует сохранить как резервный и как AAD-датчик."
    )

    doc.add_heading("1 Что ограничивает дальность", level=1)
    paragraph(
        doc,
        "На расстоянии 2 км только сферическое расхождение даёт 66,0 дБ потерь относительно уровня на 1 м. "
        "Поглощение воздуха добавляет потери, которые быстро растут с частотой. Расчёт ниже выполнен для 15 °C, "
        "относительной влажности 70 % и давления 101,325 кПа. Он не учитывает землю, растительность, здания, "
        "ветровой и температурный градиенты, поэтому является оптимистичной свободнополевой оценкой.",
    )
    doc.add_picture(str(LOSS_FIG), width=Inches(6.6))
    caption = doc.add_paragraph("Рисунок 1. Расчётные потери распространения на 2 км")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.runs[0].italic = True
    caption.runs[0].font.size = Pt(8)

    loss_rows = []
    for frequency in FREQUENCIES:
        absorption = iso9613_absorption_db_per_m(frequency) * 2000.0
        total = path_loss_db(frequency)
        received = [level - total for level in SOURCE_LEVELS]
        loss_rows.append(
            [
                str(frequency),
                f"{absorption:.1f}",
                f"{total:.1f}",
                f"{received[0]:.1f}",
                f"{received[1]:.1f}",
                f"{received[2]:.1f}",
            ]
        )
    add_table(
        doc,
        [
            "Частота, Гц",
            "Поглощение за 2 км, дБ",
            "Полные потери, дБ",
            "При 70 дБ SPL на 1 м",
            "При 85 дБ SPL на 1 м",
            "При 100 дБ SPL на 1 м",
        ],
        loss_rows,
        [1.7, 2.6, 2.2, 3.2, 3.2, 3.2],
        7.4,
    )
    paragraph(
        doc,
        "Уровни 70, 85 и 100 дБ SPL в таблице являются расчётными сценариями, а не присвоенными классам БПЛА "
        "характеристиками. В полевом исследовании малых DJI авторы указали уровни не выше 70 дБ SPL на 1 м. "
        "При таком источнике свободнополевая оценка на 2 км оказывается около 3,5 дБ SPL на 100 Гц и ниже нуля "
        "в среднем и высокочастотном диапазоне. Ни один алгоритм не восстановит сигнал, которого нет над физическим "
        "шумом микрофона и площадки.",
    )

    doc.add_heading("2 Оценка текущего T5838", level=1)
    add_table(
        doc,
        ["Параметр", "Текущее значение", "Следствие для дальности"],
        [
            ["Количество и геометрия", "4 канала, 3+1, база 0,165 м", "Почти нет пространственной направленности в полосе 100–500 Гц"],
            ["T5838 High Quality", "SNR 68 дБА, EIN 26 дБА SPL", "Собственный шум заметно выше, чем у профессиональных направленных микрофонов"],
            ["T5838 Low Power", "SNR 65 дБА, EIN 29 дБА SPL", "Режим подходит для дежурного обнаружения, но ухудшает слабосигнальный запас"],
            ["Физическая диаграмма", "Капсюль всенаправленный", "Подавление бокового шума возможно только за счёт массива и обработки"],
            ["Преимущества", "AAD 20 мкА, PDM, −40…+85 °C, высокий AOP", "Полезны для автономности, самоконтроля и резервного канала"],
        ],
        [3.2, 4.3, 9.2],
        8.0,
    )
    paragraph(
        doc,
        "Четыре одинаковых канала дают не более 6 дБ идеального выигрыша при когерентном сложении сигнала и "
        "некоррелированном собственном шуме. Это верхняя оценка второго уровня. Она не устраняет слабость первичного "
        "акустического тракта и падает при коррелированном ветровом или площадочном шуме.",
    )

    doc.add_heading("3 Кандидаты микрофонов", level=1)
    add_table(
        doc,
        ["Кандидат", "Физика и шум", "Среда и интерфейс", "Роль"],
        [
            ["TDK T5838", "68 дБА SNR, 26 дБА EIN, всенаправленный", "PDM, AAD, −40…+85 °C", "Сохранить как дежурный и резервный канал"],
            ["Infineon IM73D122", "73 дБА SNR, эквивалентно около 21 дБА EIN", "PDM, IP57 на уровне компонента, 20 Гц LFRO", "Прямая MEMS-альтернатива, примерно 5 дБ к собственному шуму"],
            ["Infineon IM73A135", "73 дБА SNR, 20 Гц LFRO, всенаправленный", "Дифференциальный аналоговый выход, IP57", "Распределённый крупноапертурный массив и резонаторные кассеты"],
            ["RØDE NTG5", "Направленный, 10 дБА self-noise, 66 мВ/Па", "P48, RF-bias, conformal coating, длина 203 мм", "Основной кандидат физического прототипа"],
            ["Sennheiser MKH 416", "Суперкардиоида и лепесток, 13 дБА, 25 мВ/Па", "P48, устойчив к влажности, −10…+70 °C", "Уличный эталон, температурный диапазон проекта не закрыт"],
        ],
        [2.8, 4.8, 5.0, 4.1],
        7.6,
    )
    paragraph(
        doc,
        "Переход от T5838 к IM73D122 или IM73A135 улучшает собственный шум примерно на 5 дБ, но сам по себе не "
        "создаёт направленность. Профессиональный NTG5 даёт около 16 дБ разницы по A-взвешенному собственному шуму "
        "относительно T5838 и сразу формирует направленную диаграмму. Это сравнение паспортных метрик, а не гарантия "
        "такого же выигрыша в ветреном поле. Для итогового выбора требуется одновременная запись на одной площадке.",
    )

    doc.add_heading("4 Физические варианты усиления", level=1)
    add_table(
        doc,
        ["Вариант", "Механизм", "Полезная полоса", "Ограничение", "Решение"],
        [
            ["8 направленных микрофонов", "Собственная секторная диаграмма и низкий шум", "Ориентировочно 300–2500 Гц", "Нужны перекрывающиеся сектора и крупная ветрозащита", "Основной прототип"],
            ["Массив 1,2–2 м", "Большая физическая база для низких частот", "Ориентировочно 100–1200 Гц", "Направленность получается после синхронного сложения", "Второй приоритет"],
            ["Парабола 0,6 м", "Фокусировка апертурой", "Практически от 500 Гц", "Узкий сектор, ветер, фазовая нестабильность", "Вспомогательный дальний канал"],
            ["Банк резонаторов", "Рост давления в узких полосах", "Расчётные центры 100–1000 Гц", "Звон, узкая полоса, чувствительность к воде и температуре", "Экспериментальный низкочастотный канал"],
            ["Гитароподобный корпус", "Набор механических и воздушных резонансов", "Зависит от формы и материала", "Неровная АЧХ и нестабильная фаза", "Не принимать без FEM и натурной калибровки"],
        ],
        [2.7, 4.1, 3.0, 4.2, 3.0],
        7.4,
    )

    doc.add_heading("5 Параболические отражатели", level=2)
    paragraph(
        doc,
        "Идеальная оценка апертурного выигрыша отражателя диаметром 0,6 м при эффективности 0,5 составляет около "
        "0 дБ на 250 Гц, 5,8 дБ на 500 Гц, 11,8 дБ на 1 кГц и 17,8 дБ на 2 кГц. Реальная конструкция будет хуже. "
        "На 2 кГц атмосферное поглощение за 2 км в принятом сценарии составляет около 17,6 дБ, поэтому высокий "
        "апертурный выигрыш в этой полосе в основном компенсирует только поглощение воздуха. На 500–1000 Гц "
        "отражатель имеет более полезный баланс, но остаётся секторным устройством.",
    )
    paragraph(
        doc,
        "Полевое исследование 2025 года показало сильную осевую дальность 18-дюймовой параболы с NTG-2, но также "
        "отметило чувствительность к порывам, узкое частотно-зависимое поле и проблемы межканальной фазы. Поэтому "
        "отражатель предлагается как вспомогательный канал после общего обнаружения, а не как единственный круговой "
        "датчик.",
    )

    doc.add_heading("6 Резонаторный банк", level=2)
    resonator_rows = []
    for frequency in (100, 160, 250, 400, 630, 1000):
        resonator_rows.append([str(frequency), f"{343.0 / (4.0 * frequency):.3f}", "Требуется поправка на открытый конец и демпфирование"])
    add_table(
        doc,
        ["Центральная частота, Гц", "Номинальная четверть волны, м", "Примечание"],
        resonator_rows,
        [4.1, 4.4, 8.0],
        8.0,
    )
    paragraph(
        doc,
        "Один резонатор не даст широкополосного усиления. Практический вариант представляет собой кассету из "
        "нескольких труб или камер с перекрывающимися резонансами и умеренным Q. Микрофон ставится в зоне максимума "
        "давления. Резонатор повышает давление относительно собственного шума капсюля только в своей полосе, но "
        "не отделяет цель от внешнего шума той же частоты и направления. Такой канал полезен для физических "
        "предфильтров гармоник, однако не должен заменять широкополосную головку.",
    )

    doc.add_heading("7 Рекомендуемый прототип", level=1)
    doc.add_picture(str(LAYOUT_FIG), width=Inches(5.7))
    caption = doc.add_paragraph("Рисунок 2. Компоновка прототипа LR-1")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.runs[0].italic = True
    caption.runs[0].font.size = Pt(8)
    add_number(doc, "Сохранить четыре штатных T5838 без изменения как AAD, резерв и контроль совместимости с EVT-PRE-20.")
    add_number(doc, "Собрать независимую головку LR-1 из восьми направленных микрофонов. Для максимума физического SNR первым кандидатом принять NTG5; NTG-2 с WS6 использовать как воспроизводимый полевой эталон из опубликованного исследования.")
    add_number(doc, "Разместить четыре канала нижнего яруса по азимутам 0°, 90°, 180° и 270° с наклоном около 10°. Верхний ярус повернуть на 45° и наклонить на 40–45°. Окончательные углы определить измеренной диаграммой, а не геометрическим эскизом.")
    add_number(doc, "Использовать кольцо диаметром около 1,2 м как начальный механический вариант. Предусмотреть сменные штанги до 2 м для проверки низкочастотной апертуры.")
    add_number(doc, "Каждый микрофон установить в полноценную корзинную ветрозащиту с внешним мехом, дренажом и развязкой от вибрации мачты. Тонкий поролон не считать достаточным для дальнего тракта.")
    add_number(doc, "Отдельно изготовить один секторный модуль: отражатель 0,6 м и кассета четвертьволновых резонаторов 100–1000 Гц. Сравнить их с направленным каналом в одном направлении, после чего решать вопрос о размножении на 360°.")
    add_number(doc, "Первую серию записывать автономным синхронным 8-канальным трактом 24 бит. Подключение к PCB-MAIN проектировать только после выбора физической головки.")

    doc.add_heading("8 Почему не следует сразу менять PCB-MIC", level=1)
    add_bullet(doc, "T5838 обеспечивает AAD и температурный диапазон, которых нет в той же форме у профессиональных shotgun-микрофонов.")
    add_bullet(doc, "Результат на 2 км определяется источником, ветром и площадкой сильнее, чем разницей между двумя MEMS на 5 дБ.")
    add_bullet(doc, "Направленная головка требует другого питания, балансного аналогового тракта или локального АЦП, защиты от импульсов и новой механики.")
    add_bullet(doc, "Парабола и резонаторы изменяют фазу и АЧХ. Их нельзя вводить в серийный пеленг без индивидуальной калибровки.")
    paragraph(
        doc,
        "Поэтому текущий EVT-PRE-20 остаётся базовым изделием. LR-1 оформляется как параллельный экспериментальный "
        "акустический модуль. После полевых данных можно выпустить PCB-MIC-LR и обновить корпус без риска задержать "
        "первую EVT-партию.",
    )

    doc.add_heading("9 Программа физической проверки", level=1)
    add_table(
        doc,
        ["Этап", "Что измерять", "Условия", "Результат"],
        [
            ["Лаборатория", "АЧХ, собственный шум, диаграмма, фазовая стабильность", "80–3000 Гц, 12 азимутов, 3 уровня", "Калибровочные файлы каждого канала"],
            ["Ветер", "Шум 20–500 Гц, импульсы, клиппинг", "0–2, 2–5 и 5–8 м/с", "Сравнение ветрозащит и механической развязки"],
            ["Дальность", "Сырой узкополосный SNR и спектр", "50, 100, 250, 500, 1000 и 2000 м", "Кривая SNR по частоте и расстоянию"],
            ["Реальные цели", "Первое обнаружение, пропуски, зависимость от курса", "Классы целей и полётный трек фиксируются заранее", "Pd и дальность отдельно по классу и погоде"],
            ["Сравнение", "T5838, NTG-2, NTG5, парабола, резонатор", "Одновременная запись на одной мачте", "Физический выигрыш без смены алгоритма"],
        ],
        [2.3, 5.0, 4.9, 4.6],
        7.7,
    )
    paragraph(
        doc,
        "Главная метрика первого этапа представляет собой разницу сырого SNR в одних и тех же временных и частотных "
        "окнах. Один алгоритм применяется ко всем каналам без отдельной подстройки. Только после этого сравниваются "
        "классификация и пеленг.",
    )

    doc.add_heading("10 Условие честной гарантии 2 км", level=1)
    paragraph(
        doc,
        "Гарантию нельзя задавать одной цифрой дальности. Перед фиксацией требования надо указать минимальный спектр "
        "источника на 1 м для каждого класса цели, допустимую скорость ветра, тип местности, температуру, влажность, "
        "высоту цели, сектор обзора, вероятность обнаружения и частоту ложных тревог. Без этих границ формулировка "
        "«2 км» остаётся пожеланием, а не проверяемым требованием.",
    )
    paragraph(
        doc,
        "Для малых электрических аппаратов с уровнем порядка 70 дБ SPL на 1 м расчёт показывает отсутствие "
        "достаточного физического запаса на 2 км даже при очень хорошем капсюле. Для более громких аппаратов и "
        "низкочастотных гармоник задача может быть достижима при тихой площадке, крупной апертуре и секторном тракте. "
        "Надёжность сети следует повышать также плотностью расстановки станций. Это уменьшает зависимость от одной "
        "двухкилометровой акустической трассы и атмосферных теневых зон.",
    )

    doc.add_heading("11 Решение для следующего шага", level=1)
    decision = doc.add_paragraph()
    decision.add_run("Предлагается утвердить не замену T5838, а изготовление физического прототипа LR-1. ").bold = True
    decision.add_run(
        "Состав первого стенда: восемь NTG5 или доступных NTG-2 с профессиональной ветрозащитой, независимый "
        "синхронный рекордер, один отражатель 0,6 м и одна кассета резонаторов. После одновременных измерений принимается "
        "решение между секторной головкой, крупноапертурным MEMS-массивом и гибридом."
    )

    sources_heading = doc.add_heading("12 Источники", level=1)
    sources_heading.paragraph_format.space_before = Pt(6)
    sources_heading.paragraph_format.space_after = Pt(3)
    references = [
        "[1] TDK InvenSense T5838 product page and DS-000383 datasheet: https://www.invensense.tdk.com/en-us/products/microphone/t5838",
        "[2] Infineon IM73D122 product data: https://www.infineon.com/part/IM73D122",
        "[3] Infineon IM73A135 datasheet: https://www.infineon.com/assets/row/public/documents/24/49/infineon-im73a135-datasheet-en.pdf",
        "[4] RØDE NTG5 specifications: https://rode.com/en-us/products/ntg5",
        "[5] Sennheiser MKH 416 specifications: https://docs.cloud.sennheiser.com/en-us/mkh-416-mkh-418-s/specs-mkh-416.html",
        "[6] Jekateryńczuk G., Piotrowski Z. Outdoor Microphone Range Tests and Spectral Analysis of UAV Acoustic Signatures for Array Development. Sensors 2025, 25, 7057: https://pmc.ncbi.nlm.nih.gov/articles/PMC12656299/",
        "[7] NASA atmospheric propagation relation and absorption discussion: https://ntrs.nasa.gov/api/citations/19750024046/downloads/19750024046.pdf",
        "[8] Meissner M. The response of Helmholtz resonator to external excitation. Archives of Acoustics: https://acoustics.ippt.pan.pl/index.php/aa/article/view/471",
    ]
    for reference in references:
        item = doc.add_paragraph(reference)
        item.paragraph_format.left_indent = Cm(0.5)
        item.paragraph_format.first_line_indent = Cm(-0.5)
        item.paragraph_format.space_after = Pt(1)
        item.paragraph_format.line_spacing = 1.0
        item.runs[0].font.size = Pt(7.5)

    doc.save(DOCX)


def build_markdown() -> None:
    rows = []
    for frequency in FREQUENCIES:
        absorption = iso9613_absorption_db_per_m(frequency) * 2000.0
        total = path_loss_db(frequency)
        received = [level - total for level in SOURCE_LEVELS]
        rows.append(
            f"| {frequency} | {absorption:.1f} | {total:.1f} | {received[0]:.1f} | {received[1]:.1f} | {received[2]:.1f} |"
        )
    markdown = f"""# Альтернативный акустический тракт Дионея для дальности от 2 км

**Документ:** {STEM}

**Ревизия:** Rev A

**Дата:** 27.10.2026

**Статус:** инженерное исследование, производственная конфигурация не изменена

## Главный вывод

Текущая четырёхмикрофонная головка T5838 подходит для малопотребляющего дежурного режима, распознавания в достижимом диапазоне и оценки направления. Она не даёт физического основания обещать распознавание малых электрических БПЛА на 2 км. Для такой дальности нужен отдельный крупноапертурный тракт с низким собственным шумом, выраженной направленностью и полноценной ветрозащитой. Существующий узел следует сохранить как резервный и как AAD-датчик.

## Расчёт потерь на 2 км

Условия: 15 °C, относительная влажность 70 %, давление 101,325 кПа. Расчёт не учитывает землю, растительность, препятствия и атмосферные теневые зоны.

| Частота, Гц | Поглощение, дБ | Полные потери, дБ | При 70 дБ SPL на 1 м | При 85 дБ SPL на 1 м | При 100 дБ SPL на 1 м |
|---:|---:|---:|---:|---:|---:|
{chr(10).join(rows)}

## Рекомендуемый прототип LR-1

1. Сохранить четыре штатных T5838 как AAD и резерв.
2. Собрать независимую восьмисекторную головку из направленных микрофонов. Первый кандидат по физическому SNR: RØDE NTG5. NTG-2 с WS6 использовать как опубликованный полевой эталон.
3. Четыре канала нижнего яруса направить по 0°, 90°, 180° и 270° с наклоном около 10°. Верхний ярус повернуть на 45° и наклонить на 40–45°.
4. Начать с кольца диаметром около 1,2 м и предусмотреть сменные штанги до 2 м.
5. Применить корзинную ветрозащиту, внешний мех, дренаж и виброразвязку.
6. Изготовить один отражатель 0,6 м и одну кассету четвертьволновых резонаторов 100–1000 Гц для одновременного сравнения.
7. Первые измерения вести автономным синхронным 8-канальным трактом. PCB-MAIN менять только после выбора физической головки.

## Честная гарантия 2 км

До фиксации требования надо задать минимальный спектр источника на 1 м для каждого класса цели, допустимую скорость ветра, местность, температуру, влажность, высоту цели, сектор обзора, вероятность обнаружения и частоту ложных тревог. Без этих условий «2 км» не является проверяемой характеристикой.

Полный анализ, сравнение микрофонов, отражателей, резонаторов и программа проверки приведены в DOCX и PDF этого документа.
"""
    MD.write_text(markdown, encoding="utf-8")


build_figures()
build_docx()
build_markdown()
print(DOCX)
print(MD)
