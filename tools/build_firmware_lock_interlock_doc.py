#!/usr/bin/env python3
"""Build and audit the EVT-PRE-20 firmware lock interlock DOCX."""

from __future__ import annotations

import importlib.util
from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/EVT_PRE_20_FIRMWARE_LOCK_INTERLOCK_REV_A.md"
OUTPUT = ROOT / "docs/EVT_PRE_20_FIRMWARE_LOCK_INTERLOCK_REV_A.docx"
BASE_BUILDER = ROOT / "outputs/doc_build/build_manual_rev_b.py"


def set_font(run, size: int) -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for key in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        fonts.set(qn(key), "Times New Roman")


def load_builder():
    spec = importlib.util.spec_from_file_location("evt_doc_builder", BASE_BUILDER)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    return module


def build() -> None:
    text = SOURCE.read_text(encoding="utf-8")
    forbidden = {"\u2014": "em dash", "\u2013": "en dash", "\u2011": "non-breaking hyphen"}
    for character, name in forbidden.items():
        if character in text:
            raise ValueError(f"source contains forbidden {name}")

    builder = load_builder()
    builder.OUT_DOCX = OUTPUT
    builder.build_docx(text)

    doc = Document(OUTPUT)
    footer = doc.sections[0].footer.paragraphs[0]
    footer.clear()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.line_spacing = 1.0
    run = footer.add_run("Дионея EVT-PRE-20. Межблокировка прошивки Rev A")
    set_font(run, 14)
    doc.core_properties.title = "Порядок разрешения блокировки прошивки станции Дионея EVT-PRE-20"
    doc.core_properties.subject = "Запрет блокировки до полного испытания таргета на реальном железе"
    doc.core_properties.author = "Проект Дионея"
    doc.core_properties.keywords = "Дионея, EVT-PRE-20, firmware lock, option bytes, EOL"
    doc.save(OUTPUT)


def audit() -> None:
    doc = Document(OUTPUT)
    errors: list[str] = []
    all_text: list[str] = []
    for index, paragraph in enumerate(doc.paragraphs, 1):
        all_text.append(paragraph.text)
        if not paragraph.text:
            continue
        if paragraph.paragraph_format.line_spacing != 1.0:
            errors.append(f"line-spacing:{index}")
        style = paragraph.style.name if paragraph.style else ""
        if style == "Title" or style.startswith("Heading"):
            if any(run.text and not run.bold for run in paragraph.runs):
                errors.append(f"heading-not-bold:{index}")
        elif re.match(r"^(?:\d+\. |• )", paragraph.text):
            pass
        else:
            indent = paragraph.paragraph_format.first_line_indent
            if indent is None or abs(indent.cm - 1.25) > 0.01:
                errors.append(f"indent:{index}")
            if paragraph.alignment != WD_ALIGN_PARAGRAPH.JUSTIFY:
                errors.append(f"alignment:{index}")
        for run in paragraph.runs:
            if run.text:
                if run.font.name != "Times New Roman":
                    errors.append(f"font:{index}:{run.font.name}")
                if run.font.size is None or round(run.font.size.pt, 1) != 14:
                    errors.append(f"size:{index}:{run.font.size}")

    for section in doc.sections:
        for paragraph in section.footer.paragraphs:
            all_text.append(paragraph.text)
            for run in paragraph.runs:
                if run.text and (run.font.name != "Times New Roman" or round(run.font.size.pt, 1) != 14):
                    errors.append("footer-font-or-size")
    combined = "\n".join(all_text)
    for character in ("\u2014", "\u2013", "\u2011"):
        if character in combined:
            errors.append(f"forbidden-dash:U+{ord(character):04X}")
    if errors:
        raise RuntimeError("DOCX format audit failed: " + ", ".join(errors[:30]))


def main() -> None:
    build()
    audit()
    print(f"wrote {OUTPUT}")
    print("DOCX format audit: PASS")


if __name__ == "__main__":
    main()
