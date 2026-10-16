from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/EVT_PRE_20_PROCUREMENT_VARIANTS_21_41_NOTICE_REV_D8.md"
OUTPUT = ROOT / "docs/EVT_PRE_20_PROCUREMENT_VARIANTS_21_41_NOTICE_REV_D8.docx"

TEXT = """# Извещение о вариантах закупки EVT PRE 20 REV D8

Документ фиксирует объём программы EVT PRE 20 и два допустимых варианта закупочной выдачи. Он предназначен для заказчика, отдела закупки, изготовителей и сборочного участка. Программа испытаний включает 41 станцию. Вариант закупки на 21 станцию является первой фазой и не изменяет состав программы или ранее выпущенный реестр серийных номеров.

## Управляющее решение

Программа EVT состоит из двух групп по 20 станций и одной стендовой станции. Зарезервированы 41 уникальный серийный номер: DIO-EVT-001 по DIO-EVT-040 и DIO-EVT-B01.

## Источники количества и BOM

Реестр вариантов находится в hardware/EVT_PROGRAM_PROCUREMENT_VARIANTS_REV_A.csv. Итоговая книга hardware/EVT_PRE_20_BOM_REV_A.xlsx содержит отдельные листы Program 21 и Program 41, а также сводную строку выбора. Машиночитаемые CSV являются источниками количества, книга XLSX предназначена для просмотра, расчёта и передачи в закупку.

## EOL и измерения

Каждая собранная станция проходит стопроцентный EOL с привязкой к собственному серийному номеру. Число 21 в строке OPEN criteria формы MFG 004 означает количество открытых критериев приёмки, а не количество станций. Окончательные числовые пределы EOL утверждаются по измерениям DIO-EVT-B01 и последующих собранных станций. Пока обязательный предел или подтверждение отсутствуют, результат остаётся HOLD или NOT RUN.

## English version

This notice fixes the EVT PRE 20 program quantity and the two permitted procurement handoffs. The EVT program contains 41 stations: two groups of 20 production stations and one bench station. Forty-one unique serial identities are reserved from DIO-EVT-001 through DIO-EVT-040 plus DIO-EVT-B01.

The customer selects exactly one option for a purchase order. Quantities for 21 and 41 are not added together. Selecting 21 is a phased first purchase and does not cancel the serial reservation for DIO-EVT-021 through DIO-EVT-040.

The controlled workbook hardware/EVT_PRE_20_BOM_REV_A.xlsx contains separate Program 21 and Program 41 worksheets and a summary selection row. Every assembled station receives 100 percent EOL against its own serial identity. The value 21 shown as OPEN criteria in the MFG 004 form is a count of unresolved acceptance criteria and is not the station quantity. Final numeric EOL limits remain open until measurements from DIO-EVT-B01 and subsequent assembled stations are reviewed and approved.

## Управление выпуском и подписи

Решение: DEC-148 от 16 октября 2026 года. Ревизия документа: D8. Статус: подготовлено для выбора варианта закупки; изготовление остаётся под действием указанных в комплекте технических ворот.
"""


def set_cell_fill(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph) -> None:
    paragraph.add_run("Страница ")
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.width = Inches(widths[idx])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_fill(cell, "1F4E78")
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(header)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
    set_repeat_table_header(table.rows[0])
    for row_idx, values in enumerate(rows):
        cells = table.add_row().cells
        for idx, value in enumerate(values):
            cell = cells[idx]
            cell.width = Inches(widths[idx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_idx % 2:
                set_cell_fill(cell, "EAF2F8")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx in (0, 1, 3) else WD_ALIGN_PARAGRAPH.LEFT
            p.add_run(value)
    doc.add_paragraph()


SOURCE.write_text(TEXT, encoding="utf-8", newline="\n")
doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.7)
section.right_margin = Inches(0.7)

normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.line_spacing = 1.08
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
for style_name, size in (("Title", 18), ("Heading 1", 15), ("Heading 2", 13)):
    style = doc.styles[style_name]
    style.font.name = "Times New Roman"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.keep_with_next = True
    style.paragraph_format.space_before = Pt(10)
    style.paragraph_format.space_after = Pt(5)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run("Извещение о вариантах закупки EVT PRE 20 REV D8")
meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = meta.add_run("Проект Дионея   Дата 16 октября 2026 года   Решение DEC-148")
r.bold = True

doc.add_paragraph(
    "Документ фиксирует объём программы EVT PRE 20 и два допустимых варианта закупочной выдачи. "
    "Он предназначен для заказчика, отдела закупки, изготовителей и сборочного участка. "
    "Программа испытаний включает 41 станцию. Вариант закупки на 21 станцию является первой "
    "фазой и не изменяет состав программы или ранее выпущенный реестр серийных номеров."
)

doc.add_heading("Управляющее решение", level=1)
doc.add_paragraph(
    "Программа EVT состоит из двух групп по 20 станций и одной стендовой станции. "
    "Зарезервированы 41 уникальный серийный номер: DIO-EVT-001 по DIO-EVT-040 и DIO-EVT-B01."
)
add_table(
    doc,
    ["Вариант", "Состав", "Серийные номера", "Резерв", "Источник количества"],
    [
        ["EVT-PROC-21", "20 + 1 = 21", "DIO-EVT-001...020 и DIO-EVT-B01", "Один пул EVT-20", "EVT_PROGRAM_1X20_PLUS_1_PROCUREMENT_REV_A.csv"],
        ["EVT-PROC-41", "20 + 20 + 1 = 41", "DIO-EVT-001...040 и DIO-EVT-B01", "Два пула EVT-20", "EVT_PROGRAM_2X20_PLUS_1_PROCUREMENT_REV_A.csv"],
    ],
    [1.0, 0.9, 1.55, 1.05, 2.5],
)
doc.add_paragraph(
    "Оба варианта доступны для выбора заказчиком. В заказе выбирают ровно один вариант. "
    "Количества вариантов 21 и 41 не складывают. Закупка 21 станции закрывает первую фазу и "
    "не отменяет резерв серийных номеров DIO-EVT-021 по DIO-EVT-040 для второй группы. "
    "Стендовая станция отдельного пула запасов не создаёт."
)

doc.add_heading("Источники количества и BOM", level=1)
doc.add_paragraph(
    "Реестр вариантов находится в hardware/EVT_PROGRAM_PROCUREMENT_VARIANTS_REV_A.csv. "
    "Итоговая книга hardware/EVT_PRE_20_BOM_REV_A.xlsx содержит отдельные листы Program 21 и "
    "Program 41, а также сводную строку выбора. Машиночитаемые CSV являются источниками количества, "
    "книга XLSX предназначена для просмотра, расчёта и передачи в закупку."
)
add_table(
    doc,
    ["Вариант", "PCB MAIN PCBA", "PCB PWR PCBA", "PCB MIC PCBA"],
    [["21", "23", "23", "86"], ["41", "45", "45", "168"]],
    [1.7, 1.7, 1.7, 1.9],
)
doc.add_paragraph(
    "Коммерческие цены остаются ориентировочными до получения предложения поставщика. "
    "Выбор варианта закупки не закрывает DFM, first article, EOL, аппаратную квалификацию или "
    "производственную блокировку прошивки."
)

doc.add_heading("EOL и измерения", level=1)
doc.add_paragraph(
    "Каждая собранная станция проходит стопроцентный EOL с привязкой к собственному серийному номеру. "
    "Число 21 в строке OPEN criteria формы MFG 004 означает количество открытых критериев приёмки, "
    "а не количество станций. Окончательные числовые пределы EOL утверждаются по измерениям "
    "DIO-EVT-B01 и последующих собранных станций. Пока обязательный предел или подтверждение "
    "отсутствуют, результат остаётся HOLD или NOT RUN."
)

doc.add_heading("English version", level=1)
doc.add_paragraph(
    "This notice fixes the EVT PRE 20 program quantity and the two permitted procurement handoffs. "
    "The EVT program contains 41 stations: two groups of 20 production stations and one bench station. "
    "Forty-one unique serial identities are reserved from DIO-EVT-001 through DIO-EVT-040 plus DIO-EVT-B01."
)
add_table(
    doc,
    ["Option", "Scope", "Serial identities", "Spare pool", "Quantity authority"],
    [
        ["EVT-PROC-21", "20 + 1 = 21", "DIO-EVT-001...020 plus DIO-EVT-B01", "One EVT-20 pool", "EVT_PROGRAM_1X20_PLUS_1_PROCUREMENT_REV_A.csv"],
        ["EVT-PROC-41", "20 + 20 + 1 = 41", "DIO-EVT-001...040 plus DIO-EVT-B01", "Two EVT-20 pools", "EVT_PROGRAM_2X20_PLUS_1_PROCUREMENT_REV_A.csv"],
    ],
    [1.0, 0.9, 1.55, 1.05, 2.5],
)
doc.add_paragraph(
    "The customer selects exactly one option for a purchase order. Quantities for 21 and 41 are not "
    "added together. Selecting 21 is a phased first purchase and does not cancel the serial reservation "
    "for DIO-EVT-021 through DIO-EVT-040. The bench station adds no separate spare pool."
)
doc.add_paragraph(
    "The controlled workbook hardware/EVT_PRE_20_BOM_REV_A.xlsx contains separate Program 21 and "
    "Program 41 worksheets and a summary selection row. Every assembled station receives 100 percent "
    "EOL against its own serial identity. The value 21 shown as OPEN criteria in the MFG 004 form is a "
    "count of unresolved acceptance criteria and is not the station quantity. Final numeric EOL limits "
    "remain open until measurements from DIO-EVT-B01 and subsequent assembled stations are reviewed and approved."
)

doc.add_heading("Управление выпуском и подписи", level=1)
doc.add_paragraph(
    "Решение DEC-148 от 16 октября 2026 года. Ревизия документа D8. Статус: подготовлено для выбора "
    "варианта закупки; изготовление остаётся под действием указанных в комплекте технических ворот."
)
for label in ("Подготовил", "Проверил", "Утвердил вариант  21  41"):
    p = doc.add_paragraph()
    p.add_run(f"{label}  ______________________________    Дата  ____________________")

for sec in doc.sections:
    footer = sec.footer.paragraphs[0]
    footer.clear()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Дионея EVT-PRE-20   Варианты закупки 21 и 41   REV D8   ")
    run.font.name = "Times New Roman"
    run.font.size = Pt(9)
    add_page_number(footer)

doc.core_properties.title = "Извещение о вариантах закупки EVT PRE 20 REV D8"
doc.core_properties.subject = "Два варианта закупки 21 и 41 при программе 41 станция"
doc.core_properties.author = "Проект Дионея"
doc.core_properties.keywords = "EVT-PRE-20, закупка, BOM, 21, 41, серийные номера"
doc.save(OUTPUT)

check = Document(OUTPUT)
all_text = "\n".join(p.text for p in check.paragraphs) + "\n" + "\n".join(
    cell.text for table in check.tables for row in table.rows for cell in row.cells
)
for phrase in (
    "DIO-EVT-001 по DIO-EVT-040 и DIO-EVT-B01",
    "EVT-PROC-21",
    "EVT-PROC-41",
    "Program 21",
    "Program 41",
    "OPEN criteria",
    "not the station quantity",
):
    if phrase not in all_text:
        raise RuntimeError(f"required phrase missing: {phrase}")
print(OUTPUT)
