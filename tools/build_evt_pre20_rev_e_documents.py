from __future__ import annotations

import importlib.util
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCE_DIR = ROOT / "docs" / "evt_pre_20_rev_e"
OUTPUT_DIR = ROOT / "outputs" / "EVT_PRE_20_REV_E_DOCUMENTS"
BASE_BUILDER = ROOT / "outputs" / "doc_build" / "build_manual_rev_b.py"

DOCUMENTS = {
    "00_Опись_и_статус_комплекта_Rev_E": (
        "Опись и статус комплекта EVT-PRE-20 Rev E",
        "Комплектность аппаратуры, программного обеспечения и эксплуатационной документации",
    ),
    "01_Техническое_задание_EVT_PRE_20_Rev_E": (
        "Техническое задание на EVT-PRE-20 Rev E",
        "Требования к партии из 40 полевых и одного стендового изделия",
    ),
    "02_Развертывание_Мухоеда_Windows11_Ubuntu2404": (
        "Развертывание сервера Мухоед",
        "Windows 11, Ubuntu Server 24.04 LTS, VPS и выделенный сервер",
    ),
    "03_Руководство_администратора_Мухоед": (
        "Руководство администратора Мухоед",
        "Учетные записи, резервное копирование, мониторинг, обновление и восстановление",
    ),
    "04_Руководство_пользователя_Мухоед": (
        "Руководство пользователя Мухоед",
        "Состояние станций, события, тревога и экран сопровождения",
    ),
    "05_Выпуск_сертификатов_и_регистрация_станций": (
        "Выпуск сертификатов и регистрация станций Дионея",
        "PKI, CSR, mTLS, pairing secret, реестр и отзыв",
    ),
    "06_Инструкция_монтажника_обновление_станции": (
        "Инструкция монтажника EVT-PRE-20",
        "Сборка, прошивка, commissioning, проверка связи и обновление",
    ),
    "07_Методика_дополнительного_дообучения_Мухоед": (
        "Методика дополнительного дообучения Мухоеда",
        "Сбор данных, контроль происхождения, обучение, приемка и выпуск модели",
    ),
    "08_Проверка_сервера_dioneya_ru": (
        "Проверка настроенного сервера dioneya.ru",
        "Боевой контур и изолированный bench на 06.10.2026",
    ),
    "09_Пошаговая_инструкция_сборщика_EVT_PRE_20": (
        "Пошаговая инструкция сборщика станции Дионея EVT-PRE-20",
        "Комплектность, механическая сборка, платы, жгуты, антенны, питание, контроль и передача на EOL",
    ),
}


def load_builder():
    spec = importlib.util.spec_from_file_location("evt_doc_builder", BASE_BUILDER)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    return module


def normalize(text: str) -> str:
    text = text.replace("\u2014", "-").replace("\u2013", "-").replace("\u2011", "-")
    return re.sub(r"[ \t]+\n", "\n", text).strip() + "\n"


def set_footer_and_metadata(path: Path, title: str, subject: str) -> None:
    doc = Document(path)
    if path.stem == "09_Пошаговая_инструкция_сборщика_EVT_PRE_20":
        for paragraph in doc.paragraphs:
            if paragraph.text.startswith("28. Финальный контрольный лист сборщика"):
                paragraph.paragraph_format.page_break_before = True
            if paragraph.text.startswith("29. Передача следующей операции"):
                paragraph.paragraph_format.page_break_before = True
    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.clear()
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.paragraph_format.line_spacing = 1.0
        footer.paragraph_format.first_line_indent = None
        run = footer.add_run("Дионея EVT-PRE-20. Комплект Rev E")
        run.font.name = "Times New Roman"
        run.font.size = Pt(14)
        fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
        for key in ("w:ascii", "w:hAnsi", "w:eastAsia"):
            fonts.set(qn(key), "Times New Roman")
    doc.core_properties.title = title
    doc.core_properties.subject = subject
    doc.core_properties.author = "Проект Дионея"
    doc.core_properties.keywords = "Дионея, EVT-PRE-20, Rev E"
    doc.save(path)


def audit(path: Path) -> dict:
    doc = Document(path)
    body = list(doc.paragraphs)
    table_paragraphs = [p for t in doc.tables for r in t.rows for c in r.cells for p in c.paragraphs]
    footer_paragraphs = [p for s in doc.sections for p in s.footer.paragraphs]
    errors: list[str] = []
    for p in body + footer_paragraphs:
        for run in p.runs:
            if not run.text:
                continue
            if run.font.name != "Times New Roman":
                errors.append(f"font:{run.font.name}:{run.text[:30]}")
            if run.font.size is None or round(run.font.size.pt, 2) != 14:
                errors.append(f"size:{run.font.size}:{run.text[:30]}")
    for p in table_paragraphs:
        for run in p.runs:
            if not run.text:
                continue
            if run.font.name != "Times New Roman":
                errors.append(f"table-font:{run.font.name}:{run.text[:30]}")
            if run.font.size is None or round(run.font.size.pt, 2) != 12:
                errors.append(f"table-size:{run.font.size}:{run.text[:30]}")
    all_text = "\n".join(p.text for p in body + table_paragraphs + footer_paragraphs)
    for dash in ("\u2014", "\u2013", "\u2011"):
        if dash in all_text:
            errors.append(f"forbidden-dash:U+{ord(dash):04X}")
    for index, p in enumerate(body, 1):
        if not p.text:
            continue
        if p.paragraph_format.line_spacing != 1.0:
            errors.append(f"line-spacing:{index}")
        style = p.style.name if p.style else ""
        if style == "Title" or style.startswith("Heading"):
            if any(run.text and not run.bold for run in p.runs):
                errors.append(f"heading-not-bold:{index}")
        elif p.text.startswith(("• ", "☐ ")) or re.match(r"^\d+\. ", p.text):
            pass
        elif p.paragraph_format.left_indent:
            pass
        else:
            indent = p.paragraph_format.first_line_indent
            if indent is None or abs(indent.cm - 1.25) > 0.01:
                errors.append(f"indent:{index}")
            if p.alignment != WD_ALIGN_PARAGRAPH.JUSTIFY:
                errors.append(f"alignment:{index}")
    return {
        "file": path.name,
        "paragraphs": len(body),
        "tables": len(doc.tables),
        "errors": errors,
    }


def main() -> None:
    builder = load_builder()
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    reports = []
    for stem, (title, subject) in DOCUMENTS.items():
        source = SOURCE_DIR / f"{stem}.md"
        if not source.exists():
            raise FileNotFoundError(source)
        markdown = normalize(source.read_text(encoding="utf-8"))
        source.write_text(markdown, encoding="utf-8", newline="\n")
        out = OUTPUT_DIR / f"{stem}.docx"
        builder.OUT_DOCX = out
        builder.build_docx(markdown)
        set_footer_and_metadata(out, title, subject)
        report = audit(out)
        reports.append(report)
        if report["errors"]:
            raise RuntimeError(json.dumps(report, ensure_ascii=False, indent=2))
    audit_path = OUTPUT_DIR / "КОНТРОЛЬ_ФОРМАТА_DOCX.json"
    audit_path.write_text(json.dumps(reports, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(reports, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
