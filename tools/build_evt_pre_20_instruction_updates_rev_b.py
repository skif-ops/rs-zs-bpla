from __future__ import annotations

import csv
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
DOCS = ROOT / "docs"
MFG = ROOT / "manufacturing"
DOCS.mkdir(parents=True, exist_ok=True)
MFG.mkdir(parents=True, exist_ok=True)

SERIALS = [f"DIO-EVT-{i:03d}" for i in range(1, 41)] + ["DIO-EVT-B01"]


HARNESS_MD = r"""# Рабочая инструкция по изготовлению и контролю жгутов Дионея EVT-PRE-20

**Документ:** EVT_PRE_20_HARNESS_MANUFACTURING_WORK_INSTRUCTION_REV_B
**Ревизия:** Rev B
**Дата:** 16.10.2026
**Комплект:** EVT-PRE-20 REV_D10
**Статус:** RELEASED FOR EVT BUILD; SUPPLIER FIRST ARTICLE REQUIRED
**Назначение:** пошаговое изготовление, контроль, регистрация и приемка жгутов для 40 полевых и одной стендовой станции

## 1. Область применения Scope

RU: Инструкция обязательна для жгутов H-MIC1…H-MIC4, H-MAIN-PWR и H-BAT-PWR. На каждую станцию требуется шесть жгутов. Для программы зарезервированы `DIO-EVT-001…040` и `DIO-EVT-B01`, всего 41 комплект. Любая замена провода, контакта, корпуса, оснастки или технологического параметра допускается только по утвержденному отклонению или ECO.

EN: This instruction governs H-MIC1…H-MIC4, H-MAIN-PWR and H-BAT-PWR. Six harnesses are required per station. The program reserves `DIO-EVT-001…040` and `DIO-EVT-B01`, 41 sets total. Any wire, terminal, housing, tooling or process-parameter substitution requires an approved deviation or ECO.

## 2. Управляющие данные Governing data

| Источник Source | Назначение Purpose | Правило Rule |
|---|---|---|
| HARNESS_LOGICAL_PINOUT_REV_A.csv | Логическая распайка Logical pinout | Номер полости и сигнал не менять Do not change cavity or signal |
| HARNESS_MANUFACTURING_SCHEDULE_REV_A.csv | Проводник за проводником Conductor schedule | Основной машинночитаемый источник Authoritative machine-readable source |
| HARNESS_EVT_LENGTH_BASIS_REV_A.csv | Длины Length basis | Длина и допуск обязательны Length and tolerance are mandatory |
| CONNECTOR_FREEZE_REV_A.csv | Корпуса и ответные части Housings and mating parts | Замена только по ECO Substitution by ECO only |
| MANUFACTURING_PARAMETER_APPROVAL_REGISTER_REV_B.csv | Параметры оснастки и процесса Tooling and process parameters | OPEN блокирует приемку OPEN blocks acceptance |

RU: При расхождении остановить работу и открыть NCR. Приоритет имеет документ с более новой утвержденной ревизией внутри одного release root. Смешивание файлов разных выпусков запрещено.

EN: Stop work and open an NCR on any conflict. The later approved revision within one release root governs. Mixing files from different releases is prohibited.

## 3. Состав одного комплекта One-station set

| Жгут Harness | Количество Qty | Длина Length | Провод Wire |
|---|---:|---:|---|
| H-MIC1…H-MIC4 | 4 | 275 ±5 mm | TE 55A0111-24, 24 AWG |
| H-MAIN-PWR | 1 | 440 ±5 mm | pins 1–6 TE 55A0111-18; pins 7–12 Alpha 3051 |
| H-BAT-PWR | 1 | 330 ±5 mm | TE 55A0111-18, 18 AWG |

## 4. Ответственность и остановка работ Responsibility and stop-work

1. RU: Оператор выполняет только действующую операционную карту и записывает фактические значения. EN: The operator follows only the current operation card and records actual values.
2. RU: Наладчик утверждает комплект оснастки и first-off образцы. EN: The setup technician approves tooling and first-off samples.
3. RU: ОТК проверяет записи, измерения и доказательства, не заполняя поля за оператора. EN: QC reviews records, measurements and evidence without completing operator fields.
4. RU: При повреждении жилы, нарушении геометрии обжима, неверной полости, неутвержденном параметре, неповеренном средстве измерения или утрате трассируемости присвоить HOLD и открыть NCR. EN: Assign HOLD and open an NCR for conductor damage, invalid crimp geometry, wrong cavity, unapproved parameter, expired calibration or lost traceability.

## 5. Подготовка рабочего места Workstation preparation

1. RU: Сверить release `EVT_PRE_20_REV_D10` и SHA-256 манифеста. EN: Verify the `EVT_PRE_20_REV_D10` release and manifest SHA-256.
2. RU: Очистить стол, включить ESD защиту, отделить партии материалов и установить контейнер для брака. EN: Clean the bench, enable ESD controls, segregate material lots and provide a reject container.
3. RU: Проверить идентификаторы резака, стриппера, обжимной оснастки, микрометра, динамометра и тестера жгутов; записать сроки калибровки. EN: Verify cutter, stripper, crimp tool, micrometer, pull tester and harness tester IDs and record calibration due dates.
4. RU: Загрузить в тестер только утвержденную распайку текущего жгута. EN: Load only the approved pinout for the current harness into the tester.

## 6. Входной контроль материалов Incoming material inspection

1. RU: Сверить MPN, производителя, партию, количество, целостность упаковки и сертификаты. EN: Verify MPN, manufacturer, lot, quantity, packaging integrity and certificates.
2. RU: Проверить соответствие сечения и наружного диаметра провода контакту и уплотнению. EN: Verify wire gauge and outer diameter against terminal and seal ranges.
3. RU: Осмотреть контакты под увеличением: деформация, коррозия, загрязнение и повреждение покрытия не допускаются. EN: Inspect terminals under magnification; deformation, corrosion, contamination and plating damage are prohibited.
4. RU: Материал со смешанной или потерянной партией изолировать. EN: Segregate mixed or untraceable lots.

## 7. Нарезка и маркировка Cut and identification

1. RU: Присвоить идентификатор `Serial-Harness-End`, например `DIO-EVT-001-H-MIC1-A`. EN: Assign a `Serial-Harness-End` identifier such as `DIO-EVT-001-H-MIC1-A`.
2. RU: Нарезать провод по HARNESS_MANUFACTURING_SCHEDULE_REV_A.csv. Измерять между datum выхода проводника из задней плоскости корпуса или наконечника. EN: Cut wire per HARNESS_MANUFACTURING_SCHEDULE_REV_A.csv. Measure between the conductor exit datum at the rear of each housing or terminal.
3. RU: Проверить цвет каждой цепи до зачистки. EN: Verify every circuit color before stripping.
4. RU: Закрепить временную маркировку так, чтобы она не попала под обжим и не закрыла дефект. EN: Apply temporary identification clear of the crimp and inspection area.

## 8. Зачистка и настройка обжима Strip and crimp setup

RU: Численные значения strip length, conductor crimp height, insulation crimp geometry и pull force берутся из действующей инструкции производителя конкретного контакта и провода либо из утвержденной карты наладки. Они не подменяются типовым значением. Поля PAR-H01…PAR-H12 в реестре параметров должны иметь статус APPROVED до приемки first article.

EN: Numeric strip length, conductor crimp height, insulation crimp geometry and pull force come from the current manufacturer instruction for the exact terminal and wire or from an approved setup sheet. Generic values are not substitutes. Parameter records PAR-H01…PAR-H12 must be APPROVED before first-article acceptance.

1. RU: Выполнить пробную зачистку; жилы не надрезаны и не отсутствуют, изоляция не расплавлена и не вытянута. EN: Make a trial strip; strands shall not be nicked or missing and insulation shall not be melted or stretched.
2. RU: Записать фактическую длину зачистки и идентификатор настройки. EN: Record actual strip length and setup ID.
3. RU: Установить терминал-специфичную матрицу, локатор и упор. EN: Install terminal-specific die, locator and stop.
4. RU: Выполнить first-off не менее чем на образцах, требуемых инструкцией поставщика и планом качества. EN: Produce first-off samples in the quantity required by the supplier instruction and quality plan.
5. RU: Измерить crimp height, выполнить разрушающий pull test, сохранить значения и фото поперечного сечения, если оно требуется планом. EN: Measure crimp height, perform destructive pull tests and retain readings and cross-section images when required by the quality plan.
6. RU: Наладчик и ОТК подписывают first-off до запуска остатка партии. EN: Setup and QC sign the first-off before the remainder of the lot is released.

## 9. Серийный обжим Production crimping

1. RU: Ввести провод до упора без подгиба жил. EN: Insert the conductor fully without folded strands.
2. RU: После каждого обжима проверить bellmouth, brush, conductor position, insulation support, seam, wings and terminal straightness. EN: After each crimp inspect bellmouth, brush, conductor position, insulation support, seam, wings and terminal straightness.
3. RU: Периодически измерять crimp height с частотой утвержденного плана контроля и после остановки, переналадки, смены рулона или оператора. EN: Measure crimp height at the approved control frequency and after stoppage, setup change, reel change or operator change.
4. RU: Не подправлять обжим плоскогубцами и не использовать контакт повторно. EN: Do not reform a crimp with pliers or reuse a terminal.

## 10. Установка контактов в корпуса Contact insertion

1. RU: До установки сверить вид со стороны mating face и нумерацию полостей. EN: Before insertion verify the mating-face view and cavity numbering.
2. RU: Вставлять контакт до фиксации первичного замка; выполнить легкую проверку удержания каждого контакта. EN: Insert until the primary lock engages and perform a light retention check on every contact.
3. RU: Установить secondary lock, если предусмотрен, и отметить его положение на фото. EN: Install the secondary lock when applicable and show its position in the evidence photo.
4. RU: Для H-BAT-PWR обеспечить правильную полярность M8 ring terminal и защитить наконечник от короткого замыкания. EN: For H-BAT-PWR preserve M8 ring-terminal polarity and protect the lug against short circuit.

## 11. Формирование жгута и защита Harness formation and protection

1. RU: Сохранить одинаковую длину H-MIC1…4; избыток укладывать сервисной петлей без перегиба. EN: Keep H-MIC1…4 equal in length and place excess in a service loop without kinking.
2. RU: Сигналы PDM_CLK, PDM_DATA и I2C вести согласно указаниям schedule, вдали от switching nodes. EN: Route PDM_CLK, PDM_DATA and I2C as specified in the schedule and away from switching nodes.
3. RU: Защита от истирания, edge guards, braid или heat-shrink не должны загрязнять acoustic ports или создавать натяг на контакте. EN: Abrasion protection, edge guards, braid or heat-shrink shall not contaminate acoustic ports or load the terminal.
4. RU: Не уменьшать радиус изгиба ниже утвержденного значения PAR-H13. EN: Do not reduce bend radius below approved parameter PAR-H13.

## 12. Контроль 100 percent electrical test

1. RU: Выполнить point-to-point continuity для каждой цепи. EN: Perform point-to-point continuity on every circuit.
2. RU: Проверить отсутствие cross-short между всеми несвязанными цепями и корпусом. EN: Verify absence of cross-shorts between all unrelated circuits and the enclosure.
3. RU: Проверить полярность H-BAT-PWR отдельным действием и независимым контролем. EN: Verify H-BAT-PWR polarity as a separate step with independent review.
4. RU: Сопротивление каждого проводника H-MIC и H-MAIN-PWR должно быть не более 0.15 ohm; H-BAT-PWR — не более 0.10 ohm. EN: Each H-MIC and H-MAIN-PWR conductor shall be no more than 0.15 ohm; H-BAT-PWR shall be no more than 0.10 ohm.
5. RU: Сохранить исходный экспорт тестера, его SHA-256 и идентификатор адаптера. EN: Retain the original tester export, its SHA-256 and adapter ID.

## 13. Размеры и внешний вид Dimensions and workmanship

1. RU: Проверить общую длину, ориентацию концов, маркировку, отсутствие острых кромок и натяга. EN: Verify overall length, end orientation, labels, absence of sharp edges and strain.
2. RU: Проверить полную посадку контактов, замков и strain relief. EN: Verify full seating of terminals, locks and strain relief.
3. RU: Критические фото: оба mating faces, first-off crimp, маркировка, общая длина и H-BAT polarity. EN: Required photos: both mating faces, first-off crimp, identification, overall length and H-BAT polarity.

## 14. First article and EVT validation

| Жгут Harness | Обязательная проверка Required validation | Решение Decision |
|---|---|---|
| H-MAIN-PWR | rail drop, heating, modem burst, I2C 100 kHz | HOLD до записи результата HOLD until recorded |
| H-BAT-PWR | 5 A load, +70 C and -40 C, polarity, ring retention, strain | HOLD до записи результата HOLD until recorded |
| H-MIC1…4 | PDM/AAD gain, phase, noise, wake, flex and strain | HOLD до записи результата HOLD until recorded |
| Все All | installed routing, clearance, edge protection, service loops | HOLD до записи результата HOLD until recorded |

## 15. Несоответствия и ремонт Nonconformance and rework

1. RU: Неверный контакт извлекать только утвержденным extractor; поврежденный контакт заменять новым. EN: Remove an incorrect terminal only with the approved extractor; replace any damaged terminal with a new one.
2. RU: После ремонта повторить визуальный контроль и полный электрический тест всего жгута. EN: After rework repeat visual inspection and the complete electrical test of the harness.
3. RU: Splice запрещен, если его нет в утвержденном чертеже и ECO. EN: A splice is prohibited unless present in the approved drawing and ECO.
4. RU: Записать NCR, исходный дефект, действие, замененные партии и повторный результат. EN: Record the NCR, original defect, action, replaced lots and retest result.

## 16. Упаковка и передача Packaging and handoff

1. RU: Закрыть mating faces защитными крышками или ESD-safe упаковкой. EN: Protect mating faces with caps or ESD-safe packaging.
2. RU: Упаковывать по серийному номеру станции, не смешивать принятые и непринятые жгуты. EN: Pack by station serial and do not mix accepted and unaccepted harnesses.
3. RU: Передать tester export, first-off record, material lots, operator/QC signatures and NCR list. EN: Hand off the tester export, first-off record, material lots, operator/QC signatures and NCR list.

## 17. Критерий PASS Gate

RU: PASS разрешен только при APPROVED всех применимых параметров, принятом first-off, полном электрическом тесте, принятых размерах и внешнем виде, наличии evidence SHA-256 и двух подписях. Пустое обязательное поле, OPEN parameter, просроченная калибровка или отсутствующий исходный лог дают HOLD.

EN: PASS requires all applicable parameters APPROVED, accepted first-off, complete electrical test, accepted dimensions and workmanship, evidence SHA-256 and two signatures. A blank mandatory field, OPEN parameter, expired calibration or missing source log produces HOLD.
"""


ASSEMBLY_MD = r"""# Руководство по сборке, первому включению и передаче на EOL Дионея EVT-PRE-20

**Документ:** EVT_PRE_20_ASSEMBLY_SETUP_INSTRUCTION_REV_B
**Ревизия:** Rev B
**Дата:** 16.10.2026
**Комплект:** EVT-PRE-20 REV_D10
**Статус:** RELEASED FOR EVT ASSEMBLY; PHYSICAL EXECUTION AND LIMIT FREEZE OPEN
**Назначение:** подробная сборка 40 полевых станций и одной стендовой станции B01 с контролем каждой операции

## 1. Основная конфигурация Baseline configuration

| Узел Assembly | Решение Baseline | Контроль Control |
|---|---|---|
| PCB-MAIN | Rev B, 8 copper layers, Altium master | DFM, stackup, impedance, DRC, AOI/X-ray and first article |
| PCB-PWR | 4 copper layers, Altium master | protection, polarity, power load and thermal test |
| PCB-MIC | 2 copper layers, Altium master, four per station | port cleanliness, identity, gain, phase and noise |
| Серийность Serials | DIO-EVT-001…040 plus DIO-EVT-B01 | 41 unique identities, no duplicates |

RU: Закупка может выполняться вариантом 21 или 41 станции, но производственный реестр и формы охватывают все 41 идентичность. Вариант закупки не изменяет плату, прошивку, распайку или критерии приемки.

EN: Procurement may use the 21- or 41-station option, while manufacturing registers and forms cover all 41 identities. Procurement quantity does not change hardware, firmware, pinout or acceptance criteria.

## 2. Условия допуска Release gate

1. RU: Сверить имя архива, SHA-256, release metadata и отсутствие файлов KiCad в производственной папке. EN: Verify archive name, SHA-256, release metadata and absence of KiCad files from the manufacturing folder.
2. RU: Использовать Altium/CAM только из вложенного архива REV_D10. EN: Use Altium/CAM only from the nested REV_D10 archive.
3. RU: Проверить подписанные отклонения, NCR и ECO до вскрытия комплекта. EN: Review signed deviations, NCRs and ECOs before opening the kit.
4. RU: Назначить serial и traveller до начала сборки. EN: Assign the serial and traveller before assembly begins.

## 3. Безопасность, ESD и инструменты Safety, ESD and tools

1. RU: Работать на проверенном ESD месте; платы хранить в защитной упаковке. EN: Use a verified ESD workstation and keep boards in protective packaging.
2. RU: Батарея, солнечная панель и внешнее питание отключены до шага первого включения. EN: Battery, solar panel and external power remain disconnected until first power.
3. RU: RF тракт не включать без антенны или 50 ohm load. EN: Do not enable the RF path without an antenna or 50 ohm load.
4. RU: Записать ID и калибровку DMM, лабораторного источника, torque tool, thermal camera, insulation tester and geometry jig. EN: Record IDs and calibration for the DMM, bench supply, torque tool, thermal camera, insulation tester and geometry jig.
5. RU: Не применять численный torque или test limit со слов оператора. EN: Do not use a numeric torque or test limit based on verbal instruction.

## 4. Входной контроль комплекта Kit inspection

1. RU: Сверить serial на traveller, корпусе, платах и labels; до provisioning допустима временная метка. EN: Match the traveller, enclosure, PCB and label serials; a temporary label is allowed before provisioning.
2. RU: Проверить PCBA records: paste/reflow profile, AOI, visual, X-ray where required, repair log and board serial. EN: Review PCBA records: paste/reflow profile, AOI, visual, required X-ray, repair log and board serial.
3. RU: Проверить четыре PCB-MIC как отдельные принятые изделия. EN: Verify four PCB-MIC assemblies as separately accepted units.
4. RU: Проверить шесть жгутов по EVT_PRE_20_HARNESS_MANUFACTURING_WORK_INSTRUCTION_REV_B. EN: Verify six harnesses under EVT_PRE_20_HARNESS_MANUFACTURING_WORK_INSTRUCTION_REV_B.
5. RU: Проверить корпус, mounts, membranes, vents, seals, fasteners, antenna/coax, battery, MPPT and labels. EN: Inspect enclosure, mounts, membranes, vents, seals, fasteners, antenna/coax, battery, MPPT and labels.

## 5. Подготовка корпуса Enclosure preparation

1. RU: Очистить посадочные места без ворса и растворителя, разрушающего пластик или мембраны. EN: Clean mounting surfaces with lint-free materials and no solvent that attacks plastics or membranes.
2. RU: Проверить отверстия, inserts, резьбу, плоскостность и отсутствие трещин. EN: Inspect holes, inserts, threads, flatness and absence of cracks.
3. RU: Выполнить сухую примерку плат и кабельных вводов; ничего не подрезать без NCR/ECO. EN: Dry-fit PCBs and glands; do not trim parts without an NCR/ECO.
4. RU: Записать lot корпуса и 3D-printed/vacuum-cast part IDs. EN: Record enclosure lot and 3D-printed/vacuum-cast part IDs.

## 6. Монтаж PCB-MAIN, антенн и коаксиалов Upper electronics module

1. RU: Установить стойки и isolators по mechanical drawing. EN: Install standoffs and isolators per the mechanical drawing.
2. RU: Убедиться, что под платой нет стружки, провода или выступающего крепежа. EN: Ensure no chips, wires or protruding fasteners remain below the board.
3. RU: Установить PCB-MAIN без изгиба. Крепеж затягивать крест-накрест в два прохода. EN: Install PCB-MAIN without bending and tighten fasteners in a cross pattern in two passes.
4. RU: Фактический torque записать; PASS возможен только при APPROVED параметре PAR-A01. EN: Record actual torque; PASS requires parameter PAR-A01 APPROVED.
5. RU: Подключить antenna/coax, проверить mating, bend radius, strain relief and antenna keepout. EN: Connect antenna/coax and verify mating, bend radius, strain relief and antenna keepout.
6. RU: Сохранить фото до закрытия модуля. EN: Retain a photo before closing the module.

## 7. Геометрия четырех микрофонов Four-microphone geometry

1. RU: Установить PCB-MIC1…4 по идентификаторам, не менять местами после calibration record. EN: Install PCB-MIC1…4 by identity and do not swap after the calibration record.
2. RU: Не касаться MEMS port инструментом, клеем, герметиком или сжатым воздухом. EN: Keep tools, adhesive, sealant and compressed air away from the MEMS port.
3. RU: Установить mounts, membranes and acoustic path in the documented orientation. EN: Install mounts, membranes and acoustic path in the documented orientation.
4. RU: Подключить H-MIC1…4, сохранив одинаковые длины; сервисные петли не пережимать. EN: Connect H-MIC1…4 at equal lengths and do not pinch service loops.
5. RU: Измерить triangle sides, MIC4 offset, heights and orientation with the approved jig; записать actual values and photo. EN: Measure triangle sides, MIC4 offset, heights and orientation with the approved jig and record actual values and a photo.
6. RU: До заморозки численных геометрических пределов B01 результат остается HOLD. EN: Until B01 freezes numeric geometry limits the result remains HOLD.

## 8. PCB-PWR, аккумулятор и нижний модуль Lower power module

1. RU: Проверить полярность battery, fuse, BMS and MPPT на обесточенном узле. EN: Verify battery, fuse, BMS and MPPT polarity while unpowered.
2. RU: Установить PCB-PWR по drawing; фактический torque записать и сопоставить с PAR-A02. EN: Install PCB-PWR per drawing, record actual torque and compare with PAR-A02.
3. RU: Подключить H-BAT-PWR: negative остается отсоединенным до финальной проверки полярности. EN: Connect H-BAT-PWR while leaving negative disconnected until final polarity review.
4. RU: M8 torque и направление наконечника принимаются только при APPROVED PAR-A03. EN: M8 torque and lug orientation are accepted only when PAR-A03 is APPROVED.
5. RU: Установить strain relief так, чтобы нагрузка не передавалась на разъем и вывод батареи. EN: Install strain relief so no load reaches the connector or battery terminal.

## 9. Межмодульные соединения Interconnect

1. RU: Подключить H-MAIN-PWR при снятом питании и проверить keying. EN: Connect H-MAIN-PWR with power removed and verify keying.
2. RU: Выполнить continuity and polarity from battery input through PCB-PWR to PCB-MAIN. EN: Check continuity and polarity from battery input through PCB-PWR to PCB-MAIN.
3. RU: Проверить отсутствие cross-short и rail-to-ground short по утвержденным пределам. EN: Verify no cross-short or rail-to-ground short against approved limits.
4. RU: Уложить силовые, RF и микрофонные линии раздельно, исключить switching-node coupling. EN: Separate power, RF and microphone routes and avoid switching-node coupling.
5. RU: Зафиксировать routing photos до закрытия корпуса. EN: Capture routing photos before closing the enclosure.

## 10. Уплотнения и закрытие Seals and closure

1. RU: Осмотреть seals, membranes, vents and glands; загрязненные или поврежденные одноразовые элементы заменить. EN: Inspect seals, membranes, vents and glands and replace contaminated or damaged single-use parts.
2. RU: Уложить seal без перекручивания, растяжения или discontinuity. EN: Seat the seal without twist, stretch or discontinuity.
3. RU: Закрывать корпус по утвержденной sequence and torque plan PAR-A04/PAR-A05. EN: Close the enclosure under approved sequence and torque plan PAR-A04/PAR-A05.
4. RU: Записать фактические torques, tool ID and calibration due date. EN: Record actual torques, tool ID and calibration due date.
5. RU: Leak test выполняется после электрической безопасности и до уличного воздействия. EN: Perform leak testing after electrical safety and before outdoor exposure.

## 11. Предпусковой контроль Pre-power checklist

1. RU: Отключить battery, solar input, USB-C and external interfaces, кроме измерительных. EN: Disconnect battery, solar input, USB-C and external interfaces except measurement connections.
2. RU: Проверить ориентацию всех полярных разъемов, fuse rating and protective devices. EN: Verify all polarized connectors, fuse rating and protective devices.
3. RU: Измерить resistance to ground каждой rail и записать meter range, settling time and value. EN: Measure each rail resistance to ground and record meter range, settling time and value.
4. RU: Подключить лабораторный источник с current limit, утвержденным в PAR-P01. EN: Connect a bench supply with the current limit approved in PAR-P01.
5. RU: Подключить thermal camera and voltage probes до подачи питания. EN: Connect the thermal camera and voltage probes before applying power.

## 12. Первое включение Initial power-up

1. RU: Подать входное напряжение по утвержденному B01 профилю; не подключать батарею одновременно. EN: Apply input voltage under the approved B01 profile and do not connect the battery simultaneously.
2. RU: Записать inrush current, steady current, 3V3, 3V8, 1V8, input voltage and timestamps. EN: Record inrush current, steady current, 3V3, 3V8, 1V8, input voltage and timestamps.
3. RU: Наблюдать thermal image и rails в течение времени PAR-P02. EN: Observe thermal image and rails for the duration in PAR-P02.
4. RU: Немедленно снять питание при current limit, smoke, odor, abnormal sound, unstable rail or abnormal heating. EN: Remove power immediately on current limit, smoke, odor, abnormal sound, unstable rail or abnormal heating.
5. RU: Не повторять включение после аварийного события без NCR disposition. EN: Do not reapply power after an abnormal event without NCR disposition.
6. RU: Сохранить source log, waveform, thermal image and SHA-256. EN: Retain source log, waveform, thermal image and SHA-256.

## 13. Подключение аккумулятора Battery connection

1. RU: Разрешается только после PASS лабораторного первого включения и независимой проверки полярности. EN: Allowed only after laboratory first-power PASS and independent polarity verification.
2. RU: Подключить positive then negative with insulated tools; disconnect in reverse order. EN: Connect positive then negative with insulated tools and disconnect in reverse order.
3. RU: Проверить отсутствие искрения, перегрева, BMS trip and reverse current. EN: Verify absence of arcing, overheating, BMS trip and reverse current.
4. RU: Выполнить charging/MPPT check по утвержденному профилю; до утверждения профиля оставить HOLD. EN: Perform charging/MPPT checks under the approved profile; keep HOLD until the profile is approved.

## 14. Передача на прошивку и provisioning Handoff

RU: Передавать станцию только с закрытыми assembly and first-power checklists, PCBA evidence, harness records, torque records, board/harness IDs, NCR list and evidence SHA-256. Firmware, Android APK, certificates, time synchronization and EOL выполняются по `EVT_PRE_20_FIRMWARE_COMMISSIONING_TEST_MANUAL_REV_D`, PKI and synchronization instructions. Candidate 2026101601 не считать final production lock до аппаратной квалификации.

EN: Transfer the station only with closed assembly and first-power checklists, PCBA evidence, harness records, torque records, board/harness IDs, NCR list and evidence SHA-256. Firmware, Android APK, certificates, time synchronization and EOL follow `EVT_PRE_20_FIRMWARE_COMMISSIONING_TEST_MANUAL_REV_D` plus PKI and synchronization instructions. Candidate 2026101601 is not the final production lock until hardware qualification closes.

## 15. EOL, ремонт и повторная проверка EOL and rework

1. RU: Полный EOL обязателен для каждой из 41 station identities. EN: Complete EOL is mandatory for every one of the 41 station identities.
2. RU: После ремонта повторить затронутую операцию, electrical safety, first power and полный regression EOL. EN: After rework repeat the affected operation, electrical safety, first power and complete regression EOL.
3. RU: Результат PASS запрещен при blank mandatory field, missing evidence, OPEN parameter, expired calibration or unresolved NCR. EN: PASS is prohibited with a blank mandatory field, missing evidence, OPEN parameter, expired calibration or unresolved NCR.

## 16. Записи и подписи Records and sign-off

| Запись Record | Обязательные поля Mandatory fields | Подпись Sign-off |
|---|---|---|
| Traveller | serial, HW IDs, operation, timestamp, result, evidence | operator plus QC |
| Torque record | joint, approved parameter, actual value, tool, calibration | operator plus QC |
| First power | setup, currents, rails, thermal, source logs, SHA-256 | operator plus reviewer |
| NCR/rework | defect, cause, action, replaced parts, retest | MRB authorization |
| Unit passport | HW/FW revisions, certificates, EOL, site handoff | QA release |

RU: Рабочие формы находятся в `EVT_PRE_20_OPERATION_CARDS_AND_CHECKLISTS_REV_B.xlsx` и соответствующих CSV. Формула или автоматически вычисленный PASS не заменяет подпись контролера.

EN: Working forms are in `EVT_PRE_20_OPERATION_CARDS_AND_CHECKLISTS_REV_B.xlsx` and the corresponding CSV files. A formula or automatically calculated PASS does not replace reviewer approval.
"""


INSTALL_MD = r"""# Инструкция по монтажу, подключению и вводу станции Дионея EVT-PRE-20

**Документ:** EVT_PRE_20_INSTALLATION_COMMISSIONING_INSTRUCTION_REV_B
**Ревизия:** Rev B
**Дата:** 16.10.2026
**Комплект:** EVT-PRE-20 REV_D10
**Статус:** RELEASED FOR SITE PREPARATION; SITE-SPECIFIC PARAMETERS AND PHYSICAL EXECUTION OPEN
**Назначение:** пошаговый монтаж, подключение, настройка, проверка и передача станции в эксплуатацию

## 1. Условия допуска Acceptance prerequisites

RU: К монтажу допускается станция с подписанным traveller, unit passport, complete EOL record, resolved NCR list and Overall PASS. Серийный номер должен входить в `DIO-EVT-001…040`; B01 устанавливается на объект только по отдельному решению. Candidate firmware 2026101601 применяется для EVT и не обозначается final production release.

EN: Installation requires a signed traveller, unit passport, complete EOL record, resolved NCR list and Overall PASS. The serial shall be within `DIO-EVT-001…040`; B01 is field-installed only by separate approval. Candidate firmware 2026101601 is for EVT and shall not be represented as final production release.

## 2. Документы на объект Site documents

1. RU: Site drawing with coordinates, elevation, orientation, foundation and service access. EN: Site drawing with coordinates, elevation, orientation, foundation and service access.
2. RU: Approved anchor and torque plan PAR-I01. EN: Approved anchor and torque plan PAR-I01.
3. RU: Grounding and lightning protection design with acceptance limit PAR-I02. EN: Grounding and lightning protection design with acceptance limit PAR-I02.
4. RU: Cable route, protective device, fuse and disconnection plan. EN: Cable route, protective device, fuse and disconnection plan.
5. RU: RF survey or confirmed coverage for both SIMs and LoRa profile. EN: RF survey or confirmed coverage for both SIMs and the LoRa profile.

## 3. Приемка площадки Site readiness

1. RU: Сверить координаты и физический site ID с заданием. EN: Match coordinates and physical site ID to the assignment.
2. RU: Проверить несущую способность, вертикальность, drainage, flood level, snow/ice and service clearance. EN: Verify load capacity, verticality, drainage, flood level, snow/ice and service clearance.
3. RU: Проверить охранную зону, доступ персонала и отсутствие постоянного препятствия микрофонам и антеннам. EN: Verify access controls and absence of permanent microphone or antenna obstructions.
4. RU: Проверить lightning protection, bonding point and power isolation. EN: Verify lightning protection, bonding point and power isolation.
5. RU: При несоответствии чертежу присвоить SITE HOLD; не адаптировать крепление на месте без approved deviation. EN: Assign SITE HOLD for drawing mismatch and do not improvise mounting without an approved deviation.

## 4. Приемка станции и распаковка Unit receipt and unpacking

1. RU: Сверить serial/QR, tamper indicators, packing list and shock/water indicators. EN: Verify serial/QR, tamper indicators, packing list and shock/water indicators.
2. RU: Осмотреть enclosure, membranes, acoustic ports, glands, antenna connectors and solar/battery interfaces. EN: Inspect enclosure, membranes, acoustic ports, glands, antenna connectors and solar/battery interfaces.
3. RU: До подъема подтвердить транспортное состояние батареи и снять только разрешенные caps. EN: Confirm battery transport state before lifting and remove only authorized caps.
4. RU: Повреждения фотографировать до распаковки и оформлять NCR. EN: Photograph damage before unpacking and open an NCR.

## 5. Механический монтаж Mechanical installation

1. RU: Использовать только крепеж, washers, isolators and anchor pattern из site drawing. EN: Use only the fasteners, washers, isolators and anchor pattern in the site drawing.
2. RU: Поднимать за предусмотренные точки; не использовать microphone arms, cable glands or antennas. EN: Lift only from designated points and never by microphone arms, cable glands or antennas.
3. RU: Установить ориентацию по drawing; записать azimuth, tilt and height. EN: Set orientation per drawing and record azimuth, tilt and height.
4. RU: Затягивать крепеж утвержденной sequence; записывать actual torque, tool ID and calibration. EN: Tighten in the approved sequence and record actual torque, tool ID and calibration.
5. RU: PASS разрешен только при APPROVED PAR-I01 и отсутствии deformation or looseness. EN: PASS requires PAR-I01 APPROVED and no deformation or looseness.

## 6. Кабели, антенны и акустика Cables, antennas and acoustics

1. RU: Сформировать drip loops ниже вводов; обеспечить service loops and strain relief. EN: Form drip loops below glands and provide service loops and strain relief.
2. RU: Соблюдать минимальные bend radii and separation от силовых трасс по утвержденному проекту. EN: Maintain approved bend radii and separation from power routes.
3. RU: Держать metalwork and cables вне antenna keepout; RF порт подключать до включения передатчика. EN: Keep metalwork and cables outside antenna keepout and connect the RF path before enabling transmission.
4. RU: Не закрывать acoustic ports, vents and membranes; проверить одинаковую ориентацию микрофонных узлов. EN: Do not obstruct acoustic ports, vents or membranes and verify microphone assembly orientation.
5. RU: Сделать обзорные и close-up photos после укладки. EN: Take overview and close-up photographs after routing.

## 7. Заземление и защита Grounding and protection

1. RU: Подключить protective bonding первым и отключать последним. EN: Connect protective bonding first and disconnect it last.
2. RU: Очистить контактную поверхность и применить approved corrosion protection. EN: Clean the bonding surface and apply approved corrosion protection.
3. RU: Измерить continuity and grounding resistance approved method; записать instrument and calibration. EN: Measure continuity and grounding resistance using the approved method and record instrument and calibration.
4. RU: До утверждения PAR-I02 результат остается HOLD. EN: Keep the result at HOLD until PAR-I02 is approved.

## 8. Подключение питания Power connection

1. RU: Убедиться, что disconnect device открыт, battery isolated and solar input covered or disconnected. EN: Verify that the disconnect device is open, the battery isolated and solar input covered or disconnected.
2. RU: Проверить polarity, open-circuit voltage, fuse rating and protective-device setting. EN: Verify polarity, open-circuit voltage, fuse rating and protective-device setting.
3. RU: Подключить ground first, then low-voltage power; активировать источник последним. EN: Connect ground first, then low-voltage power and energize the source last.
4. RU: Записать input voltage and current during start-up; compare with approved limits. EN: Record input voltage and current during start-up and compare with approved limits.
5. RU: При overcurrent, unstable rails, odor, smoke or abnormal heating отключить питание и открыть NCR. EN: De-energize and open an NCR on overcurrent, unstable rails, odor, smoke or abnormal heating.

## 9. Идентичность, сертификаты и конфигурация Identity and certificates

1. RU: Сверить QR, serial, device ID and unit passport. EN: Match QR, serial, device ID and unit passport.
2. RU: Выполнить BLE LESC pairing в контролируемом сервисном окне. EN: Perform BLE LESC pairing in the controlled service window.
3. RU: Установить site configuration без хранения секретов на телефоне. EN: Apply site configuration without retaining secrets on the phone.
4. RU: Проверить certificate fingerprint, chain, validity, CRL behavior and server identity по PKI instruction. EN: Verify certificate fingerprint, chain, validity, CRL behavior and server identity under the PKI instruction.
5. RU: Сохранить только masked identifiers and hashes в открытом отчете. EN: Store only masked identifiers and hashes in the open report.

## 10. Координаты и синхронизация Coordinates and synchronization

1. RU: Записать фактические latitude, longitude, altitude, method and accuracy evidence. EN: Record actual latitude, longitude, altitude, method and accuracy evidence.
2. RU: Выполнить configuration read-back and hash check. EN: Perform configuration read-back and hash verification.
3. RU: Проверить GNSS fix, PPS, trusted-time state, holdover transition and recovery. EN: Verify GNSS fix, PPS, trusted-time state, holdover transition and recovery.
4. RU: Численные пределы принимаются только из утвержденного EOL limit set; undefined limit gives HOLD. EN: Numeric limits come only from the approved EOL limit set; an undefined limit gives HOLD.

## 11. Связь и сервер Communications and server

1. RU: Проверить SIM1 and SIM2, APN, LTE attach, DNS and signal evidence. EN: Verify SIM1 and SIM2, APN, LTE attach, DNS and signal evidence.
2. RU: Проверить TLS, MQTT connect, heartbeat, test event, application receipt and command ACK. EN: Verify TLS, MQTT connect, heartbeat, test event, application receipt and command ACK.
3. RU: Проверить store-and-forward отключением backhaul и последующей доставкой без дублирования. EN: Verify store-and-forward by interrupting backhaul and confirming later delivery without duplication.
4. RU: Проверить LoRa profile and permitted regional configuration without radiated test outside legal limits. EN: Verify the LoRa profile and permitted regional configuration without radiated testing outside legal limits.

## 12. Акустика и функциональная проверка Acoustic and functional checks

1. RU: Проверить четыре канала, channel identity, gain, phase, noise and wake. EN: Verify four channels, channel identity, gain, phase, noise and wake.
2. RU: Осмотреть пространство вокруг microphones на отражающие или шумящие предметы и записать условия. EN: Inspect the microphone surroundings for reflective or noise-generating objects and record conditions.
3. RU: Выполнить controlled test event и подтвердить receipt в приложении. EN: Run a controlled test event and confirm receipt in the application.
4. RU: Физический site test не заменяет лабораторный EOL; оба результата сохраняются. EN: The physical site test does not replace laboratory EOL; retain both records.

## 13. Герметичность и финальный осмотр Sealing and final inspection

1. RU: Проверить glands, seals, covers, vents and fastener witness marks. EN: Inspect glands, seals, covers, vents and fastener witness marks.
2. RU: Выполнить leak test only by approved method PAR-I03; до утверждения limit результат HOLD. EN: Perform leak testing only under approved method PAR-I03; keep HOLD until the limit is approved.
3. RU: Убедиться, что service access остается возможным без разрушения трассы. EN: Verify service access remains possible without destroying routing.

## 14. Доказательства и приемка Evidence and acceptance

| Область Area | Обязательное evidence Mandatory evidence | Gate |
|---|---|---|
| Механика Mechanical | site overview, anchors, orientation, actual torques | APPROVED parameters plus reviewer |
| Электрика Electrical | polarity, ground, startup voltage/current, instrument IDs | PASS and source log |
| Связь Communications | SIM1/SIM2, TLS/MQTT, heartbeat, receipt, ACK | PASS |
| Время и GNSS Time/GNSS | fix, PPS, trust state, coordinates and hash | PASS or HOLD on open limit |
| Акустика Acoustic | four channels, identity, gain/phase/noise record | PASS or HOLD on open limit |
| Записи Records | checklist, unit passport, NCR disposition, SHA-256 | no blank mandatory field |

RU: Итоговый FIELD READY разрешен только при закрытых обязательных шагах, APPROVED site parameters, Overall PASS and installer/reviewer signatures. Любое отсутствующее evidence дает HOLD.

EN: FIELD READY requires all mandatory steps closed, APPROVED site parameters, Overall PASS and installer/reviewer signatures. Any missing evidence produces HOLD.

## 15. Отключение и демонтаж Shutdown and removal

1. RU: Остановить передачу, открыть disconnect, отключить power, then battery/solar, and remove protective bonding last. EN: Stop transmission, open the disconnect, remove power, then battery/solar, and remove protective bonding last.
2. RU: Закрыть RF and electrical connectors approved caps. EN: Cap RF and electrical connectors with approved covers.
3. RU: Зафиксировать состояние, причину демонтажа, coordinates and returned accessories. EN: Record condition, removal reason, coordinates and returned accessories.
"""


REGISTER_MD = r"""# Реестр инструкций и рабочих форм производственного комплекта EVT-PRE-20

**Документ:** EVT_PRE_20_INSTRUCTION_REGISTER_REV_B
**Ревизия:** Rev B
**Дата:** 16.10.2026
**Комплект:** EVT-PRE-20 REV_D10
**Статус:** RELEASED DOCUMENT SET; PHYSICAL RESULTS REMAIN NOT RUN UNTIL RECORDED

## 1. Назначение Purpose

RU: Реестр указывает авторитетный документ для закупки, изготовления, сборки, первого включения, прошивки, EOL, монтажа и эксплуатации. Комплект рассчитан на две группы по 20 станций и DIO-EVT-B01, всего 41 идентичность. Варианты закупки — 21 или 41. Наличие документа не закрывает физическую проверку.

EN: This register identifies the authoritative document for procurement, fabrication, assembly, first power, programming, EOL, installation and operation. The program covers two groups of 20 plus DIO-EVT-B01, 41 identities total. Procurement options are 21 or 41. Document availability does not close a physical test.

## 2. Архивы Archives

| Архив Archive | Назначение Purpose | Основное содержимое Key content |
|---|---|---|
| EVT_PRE_20_REV_D10_FULL.zip | закупка, производство, сборка, ввод | all documents, BOM, Altium/CAM, harnesses, enclosures, software, checklists |
| EVT_PRE_20_REV_D10_ALTIUM.zip | PCB/PCBA factory | PCB-MAIN 8L, PCB-PWR 4L, PCB-MIC 2L, Altium and CAM |
| EVT_PRE_20_REV_D10_TESTS.zip | program, methods, reports, Muhoed training | DOCX/PDF methods, XLSX/CSV records, validators and training source |

RU: KiCad исключен из FULL и ALTIUM, чтобы не смешивать производственный master; исходники и история сохранены в Git. Платы REV_D10 наследуют проверенное содержимое REV_D9 без изменений.

EN: KiCad is excluded from FULL and ALTIUM to prevent confusion with the manufacturing master; sources and history remain in Git. REV_D10 inherits the verified REV_D9 board content unchanged.

## 3. Управляющие инструкции Governing instructions

| Этап Stage | Документ Document | Статус Status |
|---|---|---|
| PCB/PCBA | Altium factory package and fabrication notes | vendor DFM and first article open |
| Жгуты Harnesses | EVT_PRE_20_HARNESS_MANUFACTURING_WORK_INSTRUCTION_REV_B | released; first article and supplier parameters open |
| Сборка Assembly | EVT_PRE_20_ASSEMBLY_SETUP_INSTRUCTION_REV_B | released; physical execution open |
| Прошивка и EOL Firmware/EOL | EVT_PRE_20_FIRMWARE_COMMISSIONING_TEST_MANUAL_REV_D | candidate released; hardware qualification open |
| Сертификаты PKI | EVT_PRE_20_PKI_CERTIFICATE_LIFECYCLE_INSTRUCTION_REV_A | released; device binding validation open |
| Время и координаты Time/location | EVT_PRE_20_TIME_SYNCHRONIZATION_TRUST_INSTRUCTION_REV_A | released; numeric limits open |
| Сервер Server | EVT_PRE_20_SERVER_DEPLOYMENT_BACKUP_OPERATIONS_INSTRUCTION_REV_A | released; protected host required |
| Монтаж Installation | EVT_PRE_20_INSTALLATION_COMMISSIONING_INSTRUCTION_REV_B | released; site execution open |

## 4. Рабочие формы Work forms

| Форма Form | Назначение Purpose | Охват Coverage |
|---|---|---|
| EVT_PRE_20_OPERATION_CARDS_AND_CHECKLISTS_REV_B.xlsx | единая книга операций и решений | all 41 identities |
| ASSEMBLY_ROUTE_REV_B.csv | авторитетный маршрут производства | lot plus every station |
| HARNESS_OPERATION_CARDS_REV_B.csv | порядок изготовления жгутов | six harnesses per station |
| STATION_ASSEMBLY_CHECKLIST_REV_B.csv | сборка и закрытие traveller | 41 rows |
| FIRST_POWER_CHECKLIST_REV_B.csv | первое питание и измерения | 41 rows plus repeated attempts |
| INSTALLATION_COMMISSIONING_CHECKLIST_REV_B.csv | монтаж и ввод | 41 rows |
| EVT_PRE_20_41_STATION_TRAVELLER_REGISTER_REV_B.csv | сквозной статус | 40 field plus B01 |
| MANUFACTURING_PARAMETER_APPROVAL_REGISTER_REV_B.csv | численные process limits | open/approved with source and signatures |
| EOL_RESULT_REGISTER.csv | полный EOL | 41 rows |

## 5. Последовательность применения Sequence

1. RU: Заказчик выбирает закупку 21 или 41 и фиксирует выбранный вариант. EN: The customer selects procurement quantity 21 or 41 and records it.
2. RU: Фабрика выполняет DFM, first article, PCBA inspection and records. EN: The factory completes DFM, first article, PCBA inspection and records.
3. RU: Участок жгутов закрывает first-off and 100 percent electrical test. EN: The harness area closes first-off and 100 percent electrical testing.
4. RU: Сборка выполняет traveller, mechanical assembly, torque records and initial power. EN: Assembly completes the traveller, mechanical assembly, torque records and initial power.
5. RU: Programming/EOL выполняет identities, firmware, Android, certificates, communications and full functional tests. EN: Programming/EOL completes identities, firmware, Android, certificates, communications and full functional testing.
6. RU: Монтажник закрывает site readiness, mechanical/electrical installation and commissioning. EN: The installer closes site readiness, mechanical/electrical installation and commissioning.
7. RU: QA разрешает переход только при evidence, APPROVED limits and required signatures. EN: QA authorizes transition only with evidence, APPROVED limits and required signatures.

## 6. Правило HOLD Mandatory HOLD rule

RU: Пустое обязательное поле, отсутствующий исходный лог, несовпадающий hash, OPEN parameter, просроченная калибровка, незакрытый NCR или неподписанный контроль дают HOLD. Типовое инженерное решение можно использовать для планирования и DFM, но численное значение становится приемочным только после записи источника и утверждения.

EN: A blank mandatory field, missing source log, hash mismatch, OPEN parameter, expired calibration, unresolved NCR or missing sign-off produces HOLD. A typical engineering practice may support planning and DFM, but a numeric value becomes an acceptance limit only after its source and approval are recorded.

## 7. Открытые физические ворота Open physical gates

RU: Открыты vendor DFM/first article, crimp setup limits, mechanical torques, first-power numeric envelope, leak-test limit, site anchor/grounding limits, RF/acoustic/time/GNSS numeric criteria, enclosure fit, full target pass and final firmware lock. Они закрываются данными фабрики, B01 и собранных станций. До закрытия соответствующая операция имеет HOLD, а не PASS.

EN: Vendor DFM/first article, crimp setup limits, mechanical torques, first-power numeric envelope, leak-test limit, site anchor/grounding limits, RF/acoustic/time/GNSS numeric criteria, enclosure fit, full target pass and final firmware lock remain open. Factory data, B01 and assembled units close them. Until closure the affected operation is HOLD, not PASS.
"""


DOCS_CONTENT = {
    "EVT_PRE_20_HARNESS_MANUFACTURING_WORK_INSTRUCTION_REV_B": HARNESS_MD,
    "EVT_PRE_20_ASSEMBLY_SETUP_INSTRUCTION_REV_B": ASSEMBLY_MD,
    "EVT_PRE_20_INSTALLATION_COMMISSIONING_INSTRUCTION_REV_B": INSTALL_MD,
    "EVT_PRE_20_INSTRUCTION_REGISTER_REV_B": REGISTER_MD,
}


def write_csv(path: Path, fieldnames: list[str], rows: list[dict[str, object]]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


PARAMETERS = [
    ("PAR-H01", "430300001 / Alpha 3051 strip length", "terminal and wire supplier instruction", "mm", "OPEN_SUPPLIER_FIRST_ARTICLE"),
    ("PAR-H02", "430300001 conductor crimp height", "terminal applicator setup sheet", "mm", "OPEN_SUPPLIER_FIRST_ARTICLE"),
    ("PAR-H03", "430300001 minimum pull force", "terminal and wire supplier instruction", "N", "OPEN_SUPPLIER_FIRST_ARTICLE"),
    ("PAR-H04", "430300038 / TE 55A0111-18 strip length", "terminal and wire supplier instruction", "mm", "OPEN_SUPPLIER_FIRST_ARTICLE"),
    ("PAR-H05", "430300038 conductor crimp height", "terminal applicator setup sheet", "mm", "OPEN_SUPPLIER_FIRST_ARTICLE"),
    ("PAR-H06", "430300038 minimum pull force", "terminal and wire supplier instruction", "N", "OPEN_SUPPLIER_FIRST_ARTICLE"),
    ("PAR-H07", "5040520098 / TE 55A0111-24 strip length", "terminal and wire supplier instruction", "mm", "OPEN_SUPPLIER_FIRST_ARTICLE"),
    ("PAR-H08", "5040520098 conductor crimp height", "terminal applicator setup sheet", "mm", "OPEN_SUPPLIER_FIRST_ARTICLE"),
    ("PAR-H09", "5040520098 minimum pull force", "terminal and wire supplier instruction", "N", "OPEN_SUPPLIER_FIRST_ARTICLE"),
    ("PAR-H10", "8-34114-1 / TE 55A0111-18 strip and crimp setup", "TE tooling instruction and first-off record", "controlled setup", "OPEN_SUPPLIER_FIRST_ARTICLE"),
    ("PAR-H11", "8-34114-1 minimum pull force", "TE terminal and wire instruction", "N", "OPEN_SUPPLIER_FIRST_ARTICLE"),
    ("PAR-H12", "production crimp sampling frequency", "supplier quality plan", "interval", "OPEN_SUPPLIER_FIRST_ARTICLE"),
    ("PAR-H13", "minimum installed bend radius by wire/harness", "wire and connector supplier data", "mm", "OPEN_SUPPLIER_REVIEW"),
    ("PAR-A01", "PCB-MAIN mounting torque", "approved mechanical drawing and fastener data", "N m", "OPEN_MECHANICAL_APPROVAL"),
    ("PAR-A02", "PCB-PWR mounting torque", "approved mechanical drawing and fastener data", "N m", "OPEN_MECHANICAL_APPROVAL"),
    ("PAR-A03", "battery M8 ring-terminal torque", "battery and terminal manufacturer data", "N m", "OPEN_VENDOR_APPROVAL"),
    ("PAR-A04", "enclosure fastener torque and sequence", "approved enclosure drawing and first article", "N m", "OPEN_B01_FIRST_ARTICLE"),
    ("PAR-A05", "cable-gland torque", "selected gland manufacturer data", "N m", "OPEN_VENDOR_APPROVAL"),
    ("PAR-P01", "first-power current limit and trip action", "B01 power qualification record", "A", "OPEN_B01_MEASUREMENT"),
    ("PAR-P02", "first-power dwell and thermal limit", "B01 thermal qualification record", "s / C", "OPEN_B01_MEASUREMENT"),
    ("PAR-P03", "rail resistance and voltage acceptance limits", "B01 qualification and design review", "ohm / V", "OPEN_B01_MEASUREMENT"),
    ("PAR-I01", "site anchors and tightening torque", "approved site structural design", "type / N m", "OPEN_SITE_DESIGN"),
    ("PAR-I02", "bonding and grounding acceptance limit", "approved site electrical design and local rules", "ohm", "OPEN_SITE_DESIGN"),
    ("PAR-I03", "installed enclosure leak test method and limit", "enclosure qualification and site procedure", "pressure / time / decay", "OPEN_B01_FIRST_ARTICLE"),
]


def create_parameter_csv() -> None:
    fields = ["Parameter_ID", "Parameter_RU_EN", "Required_Source", "Unit", "Approved_Value", "Status", "Approver", "Approval_Date", "Evidence_SHA256", "Notes"]
    rows = [
        {
            "Parameter_ID": pid,
            "Parameter_RU_EN": name,
            "Required_Source": source,
            "Unit": unit,
            "Approved_Value": "",
            "Status": status,
            "Approver": "",
            "Approval_Date": "",
            "Evidence_SHA256": "",
            "Notes": "PASS prohibited until APPROVED",
        }
        for pid, name, source, unit, status in PARAMETERS
    ]
    write_csv(MFG / "MANUFACTURING_PARAMETER_APPROVAL_REGISTER_REV_B.csv", fields, rows)


ASSEMBLY_OPERATIONS = [
    ("MFG-010", "QG0", "Verify REV_D10 archive and SHA-256 / Проверить архив и SHA-256", "LOT", "release index and verifier", "all hashes match", "verification log", "HOLD"),
    ("MFG-020", "QG0", "Select procurement option 21 or 41 / Выбрать вариант 21 или 41", "LOT", "procurement decision", "one option signed", "selection record", "HOLD"),
    ("MFG-030", "QG0", "Assign serial and traveller / Назначить serial и traveller", "41_IDENTITIES", "serial register", "unique DIO-EVT-001..040+B01", "traveller register", "HOLD"),
    ("MFG-040", "QG0", "Incoming inspection / Входной контроль", "ALL_MATERIAL", "BOM AVL inspection plan", "accepted or approved deviation", "inspection record", "HOLD"),
    ("MFG-050", "QG1", "Stencil paste placement and reflow / Печать пасты и оплавление", "ALL_PCBA", "released CAM PnP BOM profile", "profile accepted and traceable", "machine exports", "REWORK_OR_HOLD"),
    ("MFG-060", "QG1", "AOI visual and required X-ray / AOI визуальный и X-ray", "ALL_PCBA", "inspection plan", "no unresolved defects", "images and defect log", "REWORK_AND_REPEAT"),
    ("MFG-070", "QG1", "Bare-board rail resistance / Сопротивление шин", "ALL_PCBA", "approved limit set", "within approved limits", "measurement log", "HOLD"),
    ("MFG-080", "QG1", "PCBA first power / Первое питание плат", "PCB_MAIN_PCB_PWR", "approved B01 setup", "no trip or abnormal heating", "power and thermal log", "HOLD"),
    ("MFG-090", "QG1", "PCB-MAIN interface self-test / Самотест интерфейсов", "PCB_MAIN", "test firmware and fixture", "all fitted interfaces pass", "machine log", "REWORK_AND_FULL_RETEST"),
    ("MFG-100", "QG1", "PCB-PWR protection/load test / Нагрузочный тест", "PCB_PWR", "approved load profile", "rails protection thermal pass", "CSV waveform and thermal", "HOLD"),
    ("MFG-110", "QG1", "PCB-MIC calibration / Калибровка PCB-MIC", "PCB_MIC", "acoustic fixture", "identity gain phase noise pass", "per-board CSV", "HOLD"),
    ("MFG-120", "QG1", "Harness first-off / First-off жгутов", "HARNESS_LOT", "approved tooling parameters", "crimp height pull and visual accepted", "first-off record", "HOLD"),
    ("MFG-130", "QG1", "Build and 100 percent test harnesses / Сборка и контроль жгутов", "246_HARNESSES", "first-off accepted", "continuity polarity resistance length pass", "tester export", "REWORK_AND_FULL_RETEST"),
    ("MFG-140", "QG1", "Prepare enclosure / Подготовить корпус", "41_IDENTITIES", "accepted enclosure parts", "fit and cleanliness pass", "photos and lot IDs", "HOLD"),
    ("MFG-150", "QG1", "Install PCB-MAIN and RF / Установить PCB-MAIN и RF", "41_IDENTITIES", "accepted PCB and PAR-A01", "routing and torque pass", "traveller photos and torque", "HOLD"),
    ("MFG-160", "QG1", "Install four PCB-MIC / Установить четыре PCB-MIC", "41_IDENTITIES", "four accepted modules", "identity and port cleanliness pass", "IDs and photos", "HOLD"),
    ("MFG-170", "QG1", "Verify microphone geometry / Проверить геометрию", "41_IDENTITIES", "approved jig and limits", "all dimensions pass", "measurements and photo", "HOLD"),
    ("MFG-180", "QG1", "Install PCB-PWR and power module / Установить силовой модуль", "41_IDENTITIES", "PAR-A02/A03 approved", "polarity torque strain relief pass", "traveller and torque", "HOLD"),
    ("MFG-190", "QG1", "Connect inter-module harnesses / Соединить модули", "41_IDENTITIES", "accepted harnesses", "keying continuity polarity pass", "continuity record", "HOLD"),
    ("MFG-200", "QG1", "Install seals membranes and vents / Установить уплотнения", "41_IDENTITIES", "clean parts", "orientation and condition pass", "lot IDs and photos", "REPLACE_AND_REPEAT"),
    ("MFG-210", "QG1", "Close enclosure / Закрыть корпус", "41_IDENTITIES", "PAR-A04/A05 approved", "sequence and torques pass", "torque record", "HOLD"),
    ("MFG-220", "QG1", "Pre-power inspection / Предпусковой контроль", "41_IDENTITIES", "assembly complete", "polarity fuse resistance pass", "checklist", "HOLD"),
    ("MFG-230", "QG1", "Station initial power / Первое включение станции", "41_IDENTITIES", "PAR-P01/P02/P03 approved", "currents rails thermal pass", "source logs and SHA-256", "HOLD"),
    ("MFG-240", "QG1", "Battery and MPPT test / Проверка батареи и MPPT", "41_IDENTITIES", "bench first power pass", "polarity charging protection pass", "power log", "HOLD"),
    ("MFG-250", "QG1", "Provision identity keys and certificates / Provisioning", "41_IDENTITIES", "stable hardware identity", "unique identity and receipt", "controlled receipt", "CONTROLLED_RECOVERY"),
    ("MFG-260", "QG1", "Apply labels and QR / Нанести labels и QR", "41_IDENTITIES", "provisioning accepted", "human and QR identity match", "scan log and photo", "REPLACE_AND_RESCAN"),
    ("MFG-270", "QG2", "Run complete EOL / Выполнить полный EOL", "41_IDENTITIES", "fully assembled unit", "all mandatory tests pass", "EOL evidence bundle", "REWORK_AND_FULL_EOL"),
    ("MFG-280", "QG2", "Leak test / Проверка герметичности", "41_IDENTITIES", "PAR-I03 approved", "approved limit pass", "leak curve", "HOLD"),
    ("MFG-290", "QG2", "Final audit / Финальный аудит", "41_IDENTITIES", "EOL and leak accepted", "identity configuration records consistent", "final checklist", "HOLD"),
    ("MFG-300", "QG2", "Close traveller and passport / Закрыть traveller и паспорт", "41_IDENTITIES", "all steps closed", "no blank mandatory fields", "signed records", "HOLD"),
    ("MFG-310", "QG2", "Pack for transport / Упаковать", "40_FIELD_PLUS_B01", "unit released", "protection and pack list pass", "packing photo", "REPACK"),
    ("MFG-320", "QG2", "Lot release review / Выпуск лота", "LOT", "selected lot records closed", "zero unexplained missing units", "release minutes", "HOLD"),
]


def create_route_csv() -> None:
    fields = ["Step", "Gate", "Operation_RU_EN", "Applies_To", "Prerequisite_or_Parameters", "Acceptance", "Evidence", "Rework_or_Hold", "Status"]
    rows = [dict(zip(fields[:-1], op), Status="OPEN") for op in ASSEMBLY_OPERATIONS]
    write_csv(MFG / "ASSEMBLY_ROUTE_REV_B.csv", fields, rows)
    write_csv(MFG / "ASSEMBLY_ROUTE.csv", fields, rows)


HARNESS_OPS = [
    ("HOP-010", "Verify release and work order / Проверить выпуск и задание", "all", "manifest", "release verified", "HOLD"),
    ("HOP-020", "Verify material lots / Проверить партии материалов", "all", "incoming record", "MPN lot quantity accepted", "HOLD"),
    ("HOP-030", "Verify calibrated tools / Проверить инструмент", "all", "calibration register", "all tools valid", "HOLD"),
    ("HOP-040", "Cut and identify wires / Нарезать и маркировать", "all", "length schedule", "length color identity pass", "REWORK"),
    ("HOP-050", "Strip trial / Пробная зачистка", "by terminal-wire pair", "PAR-H01/H04/H07/H10", "no strand damage and length pass", "HOLD"),
    ("HOP-060", "Crimp setup / Наладка обжима", "by terminal-wire pair", "PAR-H02/H05/H08/H10", "tooling and first-off set", "HOLD"),
    ("HOP-070", "Pull test first-off / Pull test", "by terminal-wire pair", "PAR-H03/H06/H09/H11", "minimum pull accepted", "HOLD"),
    ("HOP-080", "QC first-off release / Выпуск first-off", "by lot", "signed setup record", "operator and QC signatures", "HOLD"),
    ("HOP-090", "Production crimp / Серийный обжим", "each end", "approved first-off", "visual and periodic crimp height pass", "REJECT_TERMINAL"),
    ("HOP-100", "Insert terminals / Вставить контакты", "each housing", "logical pinout", "cavity and retention pass", "REWORK"),
    ("HOP-110", "Install secondary locks / Установить замки", "as applicable", "connector drawing", "lock seated", "REWORK"),
    ("HOP-120", "Form and protect harness / Сформировать жгут", "each harness", "PAR-H13 and routing notes", "routing protection pass", "HOLD"),
    ("HOP-130", "100 percent continuity and cross-short / Электроконтроль", "each harness", "approved tester program", "all circuits pass", "REWORK_FULL_RETEST"),
    ("HOP-140", "Measure conductor resistance / Измерить сопротивление", "each conductor", "0.15/0.10 ohm limits", "all conductors pass", "REWORK_FULL_RETEST"),
    ("HOP-150", "Dimension and visual inspection / Размеры и внешний вид", "each harness", "drawing", "length labels workmanship pass", "REWORK"),
    ("HOP-160", "Evidence and sign-off / Доказательства и подписи", "each harness", "source logs and photos", "SHA-256 plus two signatures", "HOLD"),
    ("HOP-170", "Pack by station serial / Упаковать по serial", "each station set", "six accepted harnesses", "set complete and protected", "HOLD"),
]


def create_harness_ops_csv() -> None:
    fields = ["Operation_ID", "Operation_RU_EN", "Applies_To", "Required_Input_or_Parameter", "Acceptance", "Failure_Disposition", "Operator", "QC", "Evidence_SHA256", "Status"]
    rows = [
        {
            "Operation_ID": op[0], "Operation_RU_EN": op[1], "Applies_To": op[2],
            "Required_Input_or_Parameter": op[3], "Acceptance": op[4], "Failure_Disposition": op[5],
            "Operator": "", "QC": "", "Evidence_SHA256": "", "Status": "NOT_RUN"
        }
        for op in HARNESS_OPS
    ]
    write_csv(MFG / "HARNESS_OPERATION_CARDS_REV_B.csv", fields, rows)


def station_lot(serial: str) -> str:
    if serial == "DIO-EVT-B01":
        return "BENCH"
    number = int(serial.rsplit("-", 1)[1])
    return "LOT-1" if number <= 20 else "LOT-2"


def create_station_csvs() -> None:
    assembly_fields = [
        "Serial", "Lot", "Traveller_ID", "PCBA_Accepted", "Harness_Set_Accepted", "Enclosure_Accepted",
        "PCB_MAIN_Installed", "Four_PCB_MIC_Installed", "MIC_Geometry", "PCB_PWR_Installed", "Interconnect",
        "Seals_and_Vents", "Torque_Record", "Labels_and_QR", "Assembly_Evidence_SHA256", "NCR_ID", "Operator", "Reviewer", "Overall"
    ]
    assembly_rows = []
    for serial in SERIALS:
        row = {field: "" for field in assembly_fields}
        row.update({"Serial": serial, "Lot": station_lot(serial), "Traveller_ID": f"TRV-{serial}", "Overall": "NOT_RUN"})
        assembly_rows.append(row)
    write_csv(MFG / "STATION_ASSEMBLY_CHECKLIST_REV_B.csv", assembly_fields, assembly_rows)

    power_fields = [
        "Serial", "Attempt", "Timestamp_UTC", "PAR_P01_Status", "PAR_P02_Status", "PAR_P03_Status",
        "Polarity_Result", "Rail_Resistance_Result", "Input_V", "Current_Limit_A", "Inrush_A", "Steady_A",
        "Rail_3V3_V", "Rail_3V8_V", "Rail_1V8_V", "Thermal_Result", "Power_Log_SHA256", "Thermal_Image_SHA256",
        "NCR_ID", "Operator", "Reviewer", "Overall"
    ]
    power_rows = []
    for serial in SERIALS:
        row = {field: "" for field in power_fields}
        row.update({"Serial": serial, "Attempt": 1, "Overall": "NOT_RUN"})
        power_rows.append(row)
    write_csv(MFG / "FIRST_POWER_CHECKLIST_REV_B.csv", power_fields, power_rows)

    install_fields = [
        "Serial", "Site_ID", "Latitude", "Longitude", "Altitude_m", "PAR_I01_Status", "PAR_I02_Status", "PAR_I03_Status",
        "Site_Readiness", "Mechanical", "Grounding", "Power", "SIM1", "SIM2", "TLS_MQTT", "GNSS_Time",
        "Acoustic", "Leak_Test", "Photo_Set_SHA256", "Commissioning_Log_SHA256", "NCR_ID", "Installer", "Reviewer", "Overall"
    ]
    install_rows = []
    for serial in SERIALS:
        row = {field: "" for field in install_fields}
        row.update({"Serial": serial, "Overall": "NOT_RUN"})
        install_rows.append(row)
    write_csv(MFG / "INSTALLATION_COMMISSIONING_CHECKLIST_REV_B.csv", install_fields, install_rows)

    traveller_fields = [
        "Serial", "Lot", "Traveller_ID", "Procurement_Selected", "PCBA", "Harnesses", "Assembly", "First_Power",
        "Provisioning", "EOL", "Leak_Test", "Installation", "Unit_Passport", "Open_NCR_Count", "Evidence_Root_SHA256",
        "QA_Reviewer", "Overall"
    ]
    traveller_rows = []
    for serial in SERIALS:
        row = {field: "" for field in traveller_fields}
        row.update({"Serial": serial, "Lot": station_lot(serial), "Traveller_ID": f"TRV-{serial}", "Overall": "NOT_RUN"})
        traveller_rows.append(row)
    write_csv(MFG / "EVT_PRE_20_41_STATION_TRAVELLER_REGISTER_REV_B.csv", traveller_fields, traveller_rows)


def set_cell_fill(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_margins(cell, top=70, start=90, bottom=70, end=90) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    item = OxmlElement("w:tblHeader")
    item.set(qn("w:val"), "true")
    tr_pr.append(item)


def set_row_cant_split(row) -> None:
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def add_page_number(paragraph) -> None:
    paragraph.add_run("Страница / Page ")
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def add_run_with_emphasis(paragraph, text: str) -> None:
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(8.7)
            run.font.color.rgb = RGBColor(31, 78, 120)
        else:
            paragraph.add_run(part)


def table_widths(cols: int) -> list[float]:
    if cols == 2:
        return [2.25, 4.75]
    if cols == 3:
        return [1.65, 3.35, 2.0]
    if cols == 4:
        return [1.55, 1.0, 2.15, 2.3]
    return [7.0 / cols] * cols


def add_table(doc: Document, rows: list[list[str]]) -> None:
    headers = rows[0]
    data = rows[2:] if len(rows) > 1 and all(set(cell.strip()) <= set("-: ") for cell in rows[1]) else rows[1:]
    cols = len(headers)
    widths = table_widths(cols)
    table = doc.add_table(rows=1, cols=cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.width = Inches(widths[index])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_fill(cell, "1F4E78")
        set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(header.strip())
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(8.1)
        run.font.color.rgb = RGBColor(255, 255, 255)
    set_repeat_header(table.rows[0])
    set_row_cant_split(table.rows[0])
    for row_index, values in enumerate(data):
        row = table.add_row()
        set_row_cant_split(row)
        for index in range(cols):
            value = values[index].strip() if index < len(values) else ""
            cell = row.cells[index]
            cell.width = Inches(widths[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_index % 2:
                set_cell_fill(cell, "EDF3F8")
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(value)
            run.font.name = "Arial"
            run.font.size = Pt(8.0)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def configure(doc: Document, title: str) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.68)
    section.right_margin = Inches(0.68)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9.2)
    normal.paragraph_format.line_spacing = 1.0
    normal.paragraph_format.space_after = Pt(3.5)
    normal.paragraph_format.widow_control = True

    for name, size, before, after in (("Title", 16, 0, 9), ("Heading 1", 13.5, 10, 5), ("Heading 2", 11.5, 8, 4)):
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(20, 42, 74)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.clear()
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = header.add_run("Дионея EVT-PRE-20 | Контролируемый документ / Controlled document")
    run.font.name = "Arial"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(89, 89, 89)

    footer = section.footer.paragraphs[0]
    footer.clear()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("REV B • EVT_PRE_20_REV_D10   ")
    run.font.name = "Arial"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(89, 89, 89)
    add_page_number(footer)

    doc.core_properties.title = title
    doc.core_properties.subject = "EVT-PRE-20 bilingual controlled manufacturing instruction"
    doc.core_properties.author = "Проект Дионея"
    doc.core_properties.keywords = "EVT-PRE-20, Dioneya, manufacturing, checklist, RU, EN"


def build_docx(name: str, markdown: str) -> Path:
    source = DOCS / f"{name}.md"
    source.write_text(markdown.strip() + "\n", encoding="utf-8")
    lines = markdown.splitlines()
    doc = Document()
    configure(doc, lines[0].lstrip("# ").strip())
    index = 0
    while index < len(lines):
        line = lines[index].rstrip()
        if not line:
            index += 1
            continue
        if line.startswith("|"):
            block: list[list[str]] = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                block.append([cell.strip() for cell in lines[index].strip().strip("|").split("|")])
                index += 1
            add_table(doc, block)
            continue
        if line.startswith("# "):
            paragraph = doc.add_paragraph(style="Title")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_run_with_emphasis(paragraph, line[2:].strip())
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
        elif re.match(r"^\d+\.\s", line):
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.25)
            paragraph.paragraph_format.first_line_indent = Inches(-0.18)
            paragraph.paragraph_format.space_after = Pt(3)
            add_run_with_emphasis(paragraph, line)
        elif line.startswith("- "):
            paragraph = doc.add_paragraph(style="List Bullet")
            add_run_with_emphasis(paragraph, line[2:].strip())
        else:
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_run_with_emphasis(paragraph, line)
        index += 1
    output = DOCS / f"{name}.docx"
    doc.save(output)
    return output


def main() -> None:
    create_parameter_csv()
    create_route_csv()
    create_harness_ops_csv()
    create_station_csvs()
    for name, content in DOCS_CONTENT.items():
        output = build_docx(name, content)
        print(output)
    print(f"serial_count={len(SERIALS)} parameter_count={len(PARAMETERS)} route_count={len(ASSEMBLY_OPERATIONS)}")


if __name__ == "__main__":
    main()
